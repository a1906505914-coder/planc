# -*- coding: utf-8 -*-
"""重建 AH 公司代码.xlsx（88 家代码→公司全名映射）。
⚡⚡ 2026-08-28 修复（只配到 72 家/1070 缺失根因）：原主体来源=「客户行项目/供应商行项目」
   提取（只覆盖有往来业务的主体）→ 1070/1080/1220/1260 等 16~24 家无往来提取而缺失。
   改为【权威来源】= 科目余额表 col7（每行 `公司名 + 空格 + 4位主体代码`），
   已实测与 ah_parallel_run 三集团 88 家完全一致、零缺失。
来源：①科目余额表 col7（主来源，权威清单，实测 88/88）
      ②客户行项目/供应商行项目（兜底，仅当科目余额表缺某主体时补名）
输出：D:/底稿测试/AH/中间产物/prepared/公司代码.xlsx（sheet 公司代码，格式 |代码|名称>）
"""
import openpyxl, glob, os, re
from collections import Counter

DATA = r'D:/底稿测试/AH/数据/2026'
PREP = r'D:/底稿测试/AH/中间产物/prepared'
GROUPS = ['1010', '1357', '2468']
# 三集团 88 家权威主体清单（与 ah_parallel_run.py 一致，作为输出筛选基准）
GROUP_COMPS = {
    '1010': ['1010'],
    '1357': ['1020', '1030', '1040', '1050', '1060', '1070', '1080', '1090', '1100',
             '1170', '1220', '1240', '1250', '1260', '3010', '5020', '9020', '9030', '9070'],
    '2468': ['2010', '2020', '2030', '2040', '2050', '2070', '2080', '2090', '2100',
             '2110', '2130', '2150', '2160', '2200', '2210', '2220', '2230', '2240',
             '2250', '2260', '2270', '2280', '2290', '2300', '2310', '2330', '2340',
             '2350', '2360', '2370', '2380', '2390', '2400', '2410', '2420', '2430',
             '2450', '2460', '2470', '2480', '2490', '2500', '2510', '2520', '2530',
             '2540', '2550', '2570', '2580', '2590', '2600', '2610', '2630', '2640',
             '2650', '2660', '2680', '2690', '2700', '2720', '2730', '2740', '2760',
             '4030', '4040', '8020', '8030', '8040'],
}
ALL_88 = [c for v in GROUP_COMPS.values() for c in v]


def extract_from_tb():
    """权威来源：科目余额表 col7 = `公司名 + 空格 + 4位代码`（每主体每文件均带）。"""
    tb_dir = os.path.join(DATA, '科目余额表')
    comp2name = {}
    for fp in sorted(glob.glob(os.path.join(tb_dir, '*.xlsx'))):
        try:
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        except Exception as _ex:
            print(f'  ⚠️ 读取失败跳过：{os.path.basename(fp)}：{_ex}')
            continue
        for sh in wb.worksheets:
            it = sh.iter_rows(values_only=True)
            try:
                next(it)
            except StopIteration:
                continue
            for r in it:
                if r is None or len(r) < 8:
                    continue
                raw = str(r[7] or '').strip()
                m = re.search(r'^(.*?)\s+(\d{4})\s*$', raw)
                if m:
                    name, code = m.group(1).strip(), m.group(2)
                    comp2name.setdefault(code, name)
    return comp2name


def extract_candidates():
    """兜底：客户行项目/供应商行项目提取 代码→候选公司名（科目余额表缺失时补名用）。"""
    cand = {}   # code -> Counter(候选名)
    for folder in ['客户行项目', '供应商行项目']:
        for fp in glob.glob(os.path.join(DATA, folder, '*.xlsx')):
            try:
                wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
                ws = wb.active
                for r in ws.iter_rows(values_only=True):
                    code = str(r[2] or '')
                    txt = str(r[5] or '')
                    if not (code.isdigit() and len(code) == 4 and txt):
                        continue
                    for m in re.finditer(r'([\u4e00-\u9fff]{2,}(?:股份)?有限公司)', txt):
                        nm = m.group(1)
                        if len(nm) >= 4 and '公司' in nm:
                            cand.setdefault(code, Counter())[nm] += 1
            except Exception as _ex:
                print(f'  ⚠️ 公司名提取失败：{_ex}')
    return cand


def main():
    # ① 权威来源：科目余额表 col7
    tb_map = extract_from_tb()
    print(f'科目余额表 col7 提取: {len(tb_map)} 家')
    # ② 兜底来源：行项目候选名（仅补科目余额表缺的）
    cands = extract_candidates()
    print(f'行项目候选兜底: {len(cands)} 个代码')

    resolved = {}
    for code in ALL_88:
        if code in tb_map:
            resolved[code] = tb_map[code]
        elif code in cands:
            resolved[code] = cands[code].most_common(1)[0][0]
    # ③ 8 家非报表主体代码（仅行项目出现、无科目余额表，用户 2026-08-28 要求保留）
    EXTRA_CODES = {
        '1130': '杭州杭氧压缩机有限公司',
        '1150': '杭州杭氧钢结构设备安装有限公司',
        '1160': '杭州杭氧空分备件有限公司',
        '1210': '杭州杭氧换热设备有限公司',
        '1270': '杭氧集团股份有限公司',
        '2140': '广西杭氧金川新锐气体有限公司',
        '2190': '江西杭氧气体有限公司',
        '2670': '大连西中岛杭氧气体有限公司',
    }
    for c, nm in EXTRA_CODES.items():
        if c in cands and c not in ('1270',):
            nm = cands[c].most_common(1)[0][0]
            nm = re.sub(r'^(收回|应收|售|销|结转|部|付)', '', nm).strip()
        resolved[c] = nm
    # ⚡ 手工修正关键主体（无条件覆盖，保障 1010 集团名正确）
    OVERRIDE = {
        '1010': '杭氧集团股份有限公司',
        '2010': '江西杭氧萍钢气体有限公司',
        '2030': '河南杭氧气体有限公司',
    }
    for c, nm in OVERRIDE.items():
        if c in resolved:
            resolved[c] = nm

    # 写入 公司代码.xlsx
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '公司代码'
    for code in sorted(resolved):
        ws.append([f'|{code}|{resolved[code]}>'])
    out = os.path.join(PREP, '公司代码.xlsx')
    wb.save(out)
    print(f'已生成: {out}')
    print(f'共 {len(resolved)} 家（88 报表主体 + {len([c for c in EXTRA_CODES if c not in ALL_88])} 家保留代码）')
    missing = [c for c in ALL_88 if c not in resolved]
    if missing:
        print(f'⚠️ 88 家中缺失 {len(missing)} 家: {missing}')
    else:
        print('✅ 88 家齐全，无缺失')


if __name__ == '__main__':
    main()
