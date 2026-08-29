# -*- coding: utf-8 -*-
"""AH 全量对平核对（2026-08-17 晚，#752）：
第一层：三集团试算表「集团合计」之和 = 全公司试算表「集团合计」（逐科目）。
第二层：各集团科目底稿合计 vs 集团试算表「集团合计」（逐科目）。
"""
import glob
import os
import re
import sys

APP_DIR = os.path.dirname(os.path.abspath(__file__))
if APP_DIR not in sys.path:
    sys.path.insert(0, APP_DIR)
import paths as P

import openpyxl

PREP = os.path.join(P.DATA_DIRS['AH'], '中间产物', 'prepared')
GROUPS = {'2468': '集团A', '1010': '集团B', '1357': '集团C'}
ALL_TB = os.path.join(PREP, '自建试算表_2026.xlsx')


def load_tb(fp):
    """试算表 → {科目代码: {主体: 余额}} + 集团合计。"""
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    ws = wb['试算表'] if '试算表' in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(x or '') for x in rows[0]]
    # 主体列 = 非合计的数字列；集团合计列 = '集团合计'
    comp_ci = []
    total_ci = None
    for i, h in enumerate(hdr[2:], 2):
        if h == '集团合计':
            total_ci = i
        elif h and h not in ('', '合计', '总计'):
            comp_ci.append(i)
    out = {}
    for r in rows[1:]:
        code = str(r[0] or '').strip()
        if not code:
            continue
        comps = {}
        for ci in comp_ci:
            try:
                v = float(r[ci] or 0)
            except (TypeError, ValueError):
                v = 0.0
            comps[hdr[ci]] = v
        try:
            tot = float(r[total_ci] or 0) if total_ci is not None else sum(comps.values())
        except (TypeError, ValueError):
            tot = sum(comps.values())
        out[code] = {'name': str(r[1] or ''), 'comps': comps, 'total': tot}
    return out


def layer1():
    """三集团试算表之和 vs 全公司试算表。"""
    all_tb = load_tb(ALL_TB)
    g_tbs = {g: load_tb(os.path.join(PREP, g, '自建试算表_2026.xlsx')) for g in GROUPS}
    rows = []
    codes = sorted(set(list(all_tb) + sum([list(t) for t in g_tbs.values()], [])))
    for code in codes:
        a_tot = sum(g_tbs[g].get(code, {}).get('total', 0) for g in GROUPS)
        all_tot = all_tb.get(code, {}).get('total', 0)
        d = round(a_tot - all_tot, 2)
        if abs(d) > 0.005:
            name = all_tb.get(code, {}).get('name', '') or g_tbs['2468'].get(code, {}).get('name', '')
            rows.append((code, name, a_tot, all_tot, d))
    return rows


def layer2():
    """各集团科目底稿合计 vs 集团试算表。"""
    out_rows = []
    for g, gname in GROUPS.items():
        tb = load_tb(os.path.join(PREP, g, '自建试算表_2026.xlsx'))
        # 科目代码→名称
        code2name = {c: v['name'] for c, v in tb.items()}
        # 底稿文件（⚡ 2026-08-18 命名统一为 {科目}审计底稿_{目录名}.xlsx，原 _集团X_xxx 已改名）
        fps = glob.glob(os.path.join(PREP, g, f'*审计底稿_{g}.xlsx'))
        for fp in sorted(fps):
            base = os.path.basename(fp)
            # 科目名 = 文件名前缀
            m = re.match(r'(.+?)审计底稿_', base)
            if not m:
                continue
            subj = m.group(1)
            # 找该科目在试算表中的代码
            code = None
            for c, nm in code2name.items():
                if nm == subj or (subj in nm and len(nm) - len(subj) <= 4):
                    code = c
                    break
            # 提取底稿审定表合计
            wp_tot = extract_wp_total(fp)
            tb_tot = tb.get(code, {}).get('total', None) if code else None
            out_rows.append((g, gname, subj, code, wp_tot, tb_tot,
                             (round(wp_tot - tb_tot, 2) if wp_tot is not None and tb_tot is not None else None)))
    return out_rows


def extract_wp_total(fp):
    """底稿审定表：'公司'列 + 金额列（优先'报表数'列，其次'期末余额'）合计。"""
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    for sn in wb.sheetnames:
        if '审定' not in sn:
            continue
        ws = wb[sn]
        rows = list(ws.iter_rows(values_only=True))
        # 找表头行
        hdr_i = None
        for i, r in enumerate(rows[:4]):
            if any(str(x or '').strip() in ('公司', '主体') for x in r[:4]):
                hdr_i = i
                break
        if hdr_i is None:
            continue
        hdr = [str(x or '') for x in rows[hdr_i]]
        comp_ci = None
        amt_ci = None
        for i, h in enumerate(hdr):
            if h == '公司' and comp_ci is None:
                comp_ci = i
            if h == '报表数' or (h == '期末余额' and amt_ci is None and h != '报表数'):
                amt_ci = i
        if comp_ci is None or amt_ci is None:
            continue
        tot = 0.0
        for r in rows[hdr_i + 1:]:
            comp = str(r[comp_ci] or '').strip() if comp_ci < len(r) else ''
            # ⚡ 只统计数字主体代码行（如 2010）；跳过 应付股利/应付利息 等子目行、合计、科目行
            if not comp or not comp.isdigit():
                continue
            try:
                tot += float(r[amt_ci] or 0)
            except (TypeError, ValueError):
                pass
        return round(tot, 2)
    return None


def main():
    print('===== 第一层：三集团试算表之和 vs 全公司试算表 =====')
    r1 = layer1()
    print(f'科目差异（|差|>0.005）: {len(r1)} 个')
    for code, name, a, t, d in r1[:15]:
        print(f'  {code} {name}: 三集团{a:,.2f} vs 全公司{t:,.2f} 差{d:,.2f}')
    print()
    print('===== 第二层：集团科目底稿合计 vs 集团试算表 =====')
    r2 = layer2()
    mism = [r for r in r2 if r[6] is not None and abs(r[6]) > 1.0]
    print(f'底稿 {len(r2)} 份 | 差异>1元 {len(mism)} 份')
    for g, gname, subj, code, wp, tb, d in mism[:20]:
        print(f'  [{gname}] {subj}({code}): 底稿{wp:,.2f} vs 试算表{tb:,.2f} 差{d:,.2f}')
    # 输出结果
    return r1, r2


if __name__ == '__main__':
    main()
