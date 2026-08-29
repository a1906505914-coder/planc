# -*- coding: utf-8 -*-
"""拖拽式单文件夹终审：执行终审清理 + 补全凭证抽查对方科目。
幂等：清理重复执行无副作用。
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)


def finalize_folder(folder):
    """对单个文件夹执行终审清理 + 补全对方科目。"""
    if not folder or not os.path.isdir(folder):
        return
    # 2026-08-05 开发模式：SKIP_FINALIZE=1 时跳过整条终审链（公式浅色/合计重算/检查器耗时环节），
    # 由 regen_all/launcher 的 --skip-finalize 参数设置；开发调取数时全量 7min → 2-3min。
    if os.environ.get('SKIP_FINALIZE') == '1':
        return
    # 1) 补全凭证抽查「对方科目」（按金额精确匹配；与 builder 解耦，覆盖各程序生成的抽凭表）
    try:
        import fill_voucher_counterparty
        n = fill_voucher_counterparty.process_folder(folder)
        if n:
            print(f"[对方科目补全] {folder}: {n} 个文件")
    except Exception as ex:
        print(f"⚠️ 对方科目补全失败：{ex}")
    # 2) 小计行清理
    try:
        import cleanup_subtotal_rows
        for fn in sorted(os.listdir(folder)):
            if not (fn.endswith("_生成.xlsx") and "_bak" not in fn and not fn.startswith("~$")):
                continue
            fp = os.path.join(folder, fn)
            try:
                cleanup_subtotal_rows.cleanup_file(fp)
            except Exception as ex:
                print(f"⚠️ 重复小计清理失败 {fn}: {ex}")
    except Exception as ex:
        print(f"⚠️ 终审清理失败：{ex}")
    # 3) 合计/小计行统一重算（2026-07-31 用户要求：生成后重新加总，
    #    修正 emit 按年过滤后合计行未跟随重算的静态值——如所得税费用勾稽核对表合计）
    try:
        import recalc_totals
        n = recalc_totals.process_folder(folder)
        if n:
            print(f"[合计重算] {folder}: {n} 个单元格")
    except Exception as ex:
        print(f"⚠️ 合计重算失败：{ex}")
    # 4) 生成后自动检查（2026-07-31 用户要求：嵌套在所有小程序中——单跑 builder 与
    #    regen 共用本终审入口，生成完立即报告 合计不符/混年/抽凭空缺/空表 等问题，
    #    把「打开底稿才发现」前移到「生成时自动拦截」）
    #    2026-08-04 整合：专项检查（重分类/附注/备抵勾稽/审定数公式）已并入 audit_checker.process_folder，
    #    跑 audit_checker = 通用 + 专项全量检查，此处不再单独调用
    try:
        import audit_checker
        audit_checker.process_folder(folder)
    except Exception as ex:
        print(f"⚠️ 自动检查失败：{ex}")
    # 5) 2026-08-05 币种标注（分币种生成配套）：外币目录（泰铢/美元）的底稿在
    #    首 sheet 标题行显式标注『（币种：泰铢 THB）』，防止外币金额被误当人民币。
    #    人民币目录不标注（默认币种）；混币种目录提示先用 currency_split 拆分。
    try:
        import currency_annotate
        currency_annotate.process_folder(folder)
    except Exception as ex:
        print(f"⚠️ 币种标注失败：{ex}")
    # 6) 2026-08-03 公式单元格浅色提醒（用户要求：涉及公式部分浅色标出，便于识别联动单元格）。
    #    仅对"无填充"的公式单元格加浅黄底（合计行浅蓝等已有填充不动），幂等可重复执行。
    try:
        n = _highlight_formula_cells(folder)
        if n:
            print(f"[公式浅色] {folder}: {n} 个单元格")
    except Exception as ex:
        print(f"⚠️ 公式浅色失败：{ex}")


def _highlight_formula_cells(folder):
    """扫描文件夹全部 *_生成.xlsx，对无填充的公式单元格（值以 = 开头）加浅黄底。返回标记数。"""
    import openpyxl
    from openpyxl.styles import PatternFill
    HL = PatternFill('solid', fgColor='FFF2CC')  # 浅黄：黑字可读
    n = 0
    for fn in sorted(os.listdir(folder)):
        if not (fn.endswith('_生成.xlsx') and not fn.startswith('~$') and '_bak' not in fn):
            continue
        fp = os.path.join(folder, fn)
        try:
            wb = openpyxl.load_workbook(fp)
        except Exception:
            continue
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if not (isinstance(cell.value, str) and cell.value.startswith('=')):
                            continue
                        f = cell.fill
                        if f is None or f.patternType is None or f.patternType == 'none':
                            cell.fill = HL
                            n += 1
            wb.save(fp)
        except Exception:
            pass
        finally:
            wb.close()
    return n
