# -*- coding: utf-8 -*-
"""审计底稿按年输出公共模块（2026-07-27）。
让各 detail 程序在构建完"合并逻辑"的 wb 后，直接调用 emit_per_year 生成
<name>审计底稿_2025_生成.xlsx / _2026_生成.xlsx，由各 detail 程序内置调用，不再需要外部拆分器后处理。

拆分规则（基于 sheet 名称中的年份标记）：
 - sheet 名含 '2025'(不含'2026') → 仅入 2025 文件
 - sheet 名含 '2026'(不含'2025') → 仅入 2026 文件
 - 同时含两年份 / 都不含年份标记 → 视为两年对比，两份都保留
 - 完全无年份专属 sheet → 若传入了 all_years 则强制按年拆分；否则保留单份
"""
import os
import sys
import glob
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

KEEP_KEYWORDS = ["交易性金融资产", "长期股权投资减值准备", "长期股权投资", "商誉",
                 "应付票据", "租赁负债", "专项应付款", "预计负债", "递延所得税负债"]


def _clean_empty_default(wb):
    """删除构建过程中产生的空默认 Sheet。"""
    if "Sheet" in wb.sheetnames:
        s = wb["Sheet"]
        if s.max_row <= 1 and s.max_column <= 1 and s.cell(1, 1).value is None:
            del wb["Sheet"]


def emit_per_year(wb_or_path, base_name, out_dir, remove_src=True, all_years=None, skip_filter_sheets=()):
    """把构建好的 wb（或已保存的合并稿路径）按 sheet 年份标记拆成 2025/2026 文件并保存（替代 wb.save）。
    返回生成的文件路径列表。

    支持两种输入：
      - 传入 openpyxl.Workbook 对象：直接拆分（内部写临时稿再删另一年 sheet）。
      - 传入路径字符串：先加载该合并稿，拆分后默认删除原合并稿（remove_src=True），
        便于"先合并产出→再审定表注入→最后按年拆分"的流程。

    all_years: 当 sheet 名不含年份标记时，若传入了此参数则强制按列表中的年份拆分（每份含全部 sheet）；
              不传时保留原行为（纯两年对比稿存为单文件）。
    2026-07-29 新增 all_years 参数以支持职工薪酬/营业收入/银行存款/借款/关联交易等无年份标记底稿的强制拆分。
    skip_filter_sheets: 拆分后按年行级过滤时跳过的 sheet 名集合（2026-07-31 新增）——
                        「期间对比」等刻意保留跨年段的表（银行存款期间对比须含 2025+2026 双段+总计），
                        避免 _filter_sheet_rows_by_year 把非目标年段删掉。
    """
    if not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)
    loaded_here = False
    if isinstance(wb_or_path, str):
        src_path = wb_or_path
        wb = openpyxl.load_workbook(src_path)
        loaded_here = True
    else:
        wb = wb_or_path
    _clean_empty_default(wb)
    sheets = list(wb.sheetnames)
    y25_only = [s for s in sheets if "2025" in s and "2026" not in s]
    y26_only = [s for s in sheets if "2026" in s and "2025" not in s]
    outs = []
    if not y25_only and not y26_only:
        if all_years:
            # ⚡⚡ 2026-08-25 性能：各年份文件内容完全相同（sheet 名无年份标记且不按年过滤）
            #   → 只 save 一次 + shutil.copyfile（原对每年度 load tmp + save，85MB 底稿 3 次全量 IO）。
            import shutil as _sh
            first = os.path.join(out_dir, f"{base_name}_{all_years[0]}_生成.xlsx")
            wb.save(first)
            outs.append(first)
            for _y in all_years[1:]:
                _o = os.path.join(out_dir, f"{base_name}_{_y}_生成.xlsx")
                try:
                    _sh.copyfile(first, _o)
                except BaseException:
                    wb.save(_o)
                outs.append(_o)
        else:
            final = os.path.join(out_dir, f"{base_name}_生成.xlsx")
            wb.save(final)
            outs.append(final)
    else:
        # ⚡⚡ 2026-08-25 性能重构：原 1×save(tmp)+2×(load+save 拆分)+2×(load+save 行级过滤)=5 次全量 IO，
        #   对 85MB 大底稿（XBJ ORA）每次 IO 数十秒 → 5× 反复写入是"2 小时单科目"主因之一。
        #   优化：
        #   ① 单年度（仅 2025 或仅 2026）：内存中直接行级过滤 + save（1 次 IO，无 tmp 落盘与重载）
        #   ② 双年度：tmp 仍落盘 1 次，但【拆分与行级过滤合并】——load 一次同时删 sheet+过滤再 save
        #      （3 次 IO），行级过滤不再对已落盘文件二次 load。
        try:
            from audit_year_io import _filter_sheet_rows_by_year as _filt
        except Exception:
            _filt = None
        if y25_only and not y26_only:
            _out = os.path.join(out_dir, f"{base_name}_2025_生成.xlsx")
            if _filt:
                try:
                    _filt(wb, '2025', skip_sheets=skip_filter_sheets)
                except Exception:
                    pass
            wb.save(_out)
            outs.append(_out)
        elif y26_only and not y25_only:
            _out = os.path.join(out_dir, f"{base_name}_2026_生成.xlsx")
            if _filt:
                try:
                    _filt(wb, '2026', skip_sheets=skip_filter_sheets)
                except Exception:
                    pass
            wb.save(_out)
            outs.append(_out)
        else:
            tmp = os.path.join(out_dir, f".{base_name}__tmp_生成.xlsx")
            wb.save(tmp)
            try:
                if y25_only:
                    w = openpyxl.load_workbook(tmp)
                    for s in y26_only:
                        if s in w.sheetnames:
                            del w[s]
                    if _filt:
                        try:
                            _filt(w, '2025', skip_sheets=skip_filter_sheets)
                        except Exception:
                            pass
                    out = os.path.join(out_dir, f"{base_name}_2025_生成.xlsx")
                    w.save(out)
                    w.close()
                    outs.append(out)
                if y26_only:
                    w2 = openpyxl.load_workbook(tmp)
                    for s in y25_only:
                        if s in w2.sheetnames:
                            del w2[s]
                    if _filt:
                        try:
                            _filt(w2, '2026', skip_sheets=skip_filter_sheets)
                        except Exception:
                            pass
                    out = os.path.join(out_dir, f"{base_name}_2026_生成.xlsx")
                    w2.save(out)
                    w2.close()
                    outs.append(out)
            finally:
                if os.path.exists(tmp):
                    try:
                        os.remove(tmp)
                    except BaseException:
                        pass
    if loaded_here:
        wb.close()
        if remove_src and os.path.exists(src_path) and src_path not in outs:
            try:
                os.remove(src_path)
            except BaseException:
                print(f'  ⚠️ 合并稿删除失败（沙箱拦截，可手动清理）：{os.path.basename(src_path)}')
    return outs


def cleanup_group(folder):
    """清理【集团】* 文件：9 个其他类关键词保留并去前缀，其余删除。"""
    for gpath in glob.glob(os.path.join(folder, "【集团】*_生成.xlsx")):
        gfn = os.path.basename(gpath)
        inner = gfn[len("【集团】"):]
        keep = any(kw in inner for kw in KEEP_KEYWORDS)
        if keep:
            target = os.path.join(folder, inner)
            if os.path.exists(target):
                try:
                    os.remove(target)
                except BaseException:
                    pass
            os.rename(gpath, target)
        else:
            try:
                os.remove(gpath)
            except BaseException:
                pass


def rebuild_audit_per_year(outs, data_dir, audit_items):
    """对 emit_per_year 拆分后的每份文件，按该年 target_year 重新生成审定表。
    
    outs: emit_per_year 返回的文件路径列表
    data_dir: 数据文件夹路径
    audit_items: list[dict]，每个元素传给 add_audit_summary_sheets 的 item
                 自动从文件名提取年份作为 target_year
    """
    import re as _re
    from audit_common import add_audit_summary_sheets as _rebuild_audit
    from audit_common import discover_entities as _de, read_tb_full as _rtb
    # 预发现实体和TB，避免每个文件重复读取
    _entities = _de(data_dir)
    _tb = _rtb(data_dir, _entities)
    for fp in outs:
        base = os.path.basename(fp)
        m = _re.search(r'_(\d{4})_', base)
        if not m:
            continue
        y_target = m.group(1)
        try:
            wb = openpyxl.load_workbook(fp)
            # 删除已有的审定表（如果有，重新生成）
            for item in audit_items:
                title = item.get('title', '审定表')
                if title in wb.sheetnames:
                    del wb[title]
            # 重新生成审定表，指定 target_year
            _rebuild_audit(wb, data_dir, audit_items, target_year=y_target,
                          tb_full=_tb, entities=_entities)
            # 审定表置于首位
            sheets = wb._sheets
            for item in audit_items:
                title = item.get('title', '审定表')
                idx = next((i for i, s in enumerate(sheets) if s.title == title), None)
                if idx is not None and idx > 0:
                    sheets.insert(0, sheets.pop(idx))
            # 过滤非年份标记 sheet 的行级数据（如期间对比只保留目标年）
            _filter_sheet_rows_by_year(wb, y_target)
            from audit_shell import finalize_workbook as _fw
            _fw(wb)
            wb.save(fp)
            wb.close()
        except Exception as ex:
            print(f'  ⚠️ 重生成审定表失败 {base}：{ex}')


def _rewrite_voucher_total(ws):
    """按年过滤删除行后，重写凭证抽查表合计行的 SUM 公式（引用范围随实际数据行数调整），
    避免"2026 文件合计显示全部年度 61 笔"失真（2026-07-31 修复）。"""
    import openpyxl as _oxl
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 1).value == '(合计)':
            last = r - 1
            if last > 2:
                ws.cell(r, 8, f'=SUM(H3:H{last})')
                ws.cell(r, 9, f'=SUM(I3:I{last})')
            break


def _del_empty_recon_sheet(wb, sn, ws):
    """核对/勾稽类 sheet 无任何金额数据单元格 → 删除该 sheet（2026-07-31 用户要求"勾稽核对无金额不产生"）。
    注意：只删"核对/勾稽"类 sheet（往来款对方科目核对/递延所得税核对等），
    不误伤审定表/明细表/凭证抽查表（凭证抽查表即使无数据也保留表头+合计供参考）。"""
    if '核对' not in sn and '勾稽' not in sn:
        return
    for _r in range(1, ws.max_row + 1):
        for _c in range(1, min(ws.max_column + 1, 13)):
            _v = ws.cell(_r, _c).value
            if isinstance(_v, (int, float)) and abs(_v) > 0.005:
                return  # 有金额 → 保留
    del wb[sn]


def _filter_sheet_rows_by_year(wb, y_target, skip_sheets=()):
    """过滤工作簿中无年份标记 sheet 的行级数据，只保留目标年份的行。
    目前支持：期间对比（会计期间列）、对方科目核对（年度列）等。
    skip_sheets: 明确跳过的 sheet 名（保留跨年段的表，如银行存款「期间对比」）。"""
    import re as _re
    for sn in list(wb.sheetnames):
        if sn in skip_sheets:
            continue  # 刻意保留跨年段（2026-07-31）
        if _re.search(r'202[0-9]', sn):
            continue  # 已有年份标记由 emit_per_year 处理
        ws = wb[sn]
        yr_col = None
        # 年份列定位须【精确匹配】表头词（年度/年份/会计期间/会计年度）或 4 位年份单元格值，
        # 不能用"包含"——如"所得税费用勾稽核对表（…以前年度…）"标题含'年度'会被误判为年份列，
        # 导致标题下所有行（含表头/说明/数据）被当作非目标年删除，只剩标题行。
        for r in range(1, min(ws.max_row + 1, 6)):
            for c in range(1, min(ws.max_column + 1, 12)):
                v = ws.cell(r, c).value
                if v:
                    vs = str(v).strip()
                    # '期间' 列名（往来款对方科目核对用，值=年份）也识别；2026-07-31 新增
                    if vs in ('会计期间', '会计年度', '年度', '年份', '期间') or _re.fullmatch(r'(?:19|20)\d{2}', vs):
                        yr_col = c
                        break
            if yr_col:
                break
        if yr_col is None:
            # 无年份列 → 尝试【日期列】（凭证抽查表等无年份标记 sheet）：按日期前 4 位年份过滤，
            # 解决"凭证抽查表被 emit 复制到各年份文件后混入他年凭证"（如 2026 文件混 2025 抽凭）。
            date_col = None
            for r in range(1, min(ws.max_row + 1, 4)):
                for c in range(1, min(ws.max_column + 1, 16)):
                    v = ws.cell(r, c).value
                    if v and str(v).strip() in ('日期', '凭证日期'):
                        date_col = c
                        break
                if date_col:
                    break
            if date_col is None:
                continue
            for r in range(ws.max_row, 1, -1):
                v = ws.cell(r, date_col).value
                if v:
                    _dm = _re.search(r'(?:19|20)\d{2}', str(v))
                    if _dm and _dm.group(0) != str(y_target):
                        ws.delete_rows(r)
            _rewrite_voucher_total(ws)
            _del_empty_recon_sheet(wb, sn, ws)
            continue
        for r in range(ws.max_row, 1, -1):
            v = ws.cell(r, yr_col).value
            if v is not None:
                vs = str(v).strip()
                # 跳过表头行（含年份/年度关键词的不属于数据行）
                if vs in ('年份', '年度', '会计年度', '会计期间', '期间'):
                    continue
                # 仅当值为 4 位年份且 ≠ 目标年才删除（2026-07-31：避免 '1-3月' 等
                # 非年份期间值被当作目标年不匹配而误删）
                if _re.fullmatch(r'(?:19|20)\d{2}', vs) and vs != str(y_target):
                    ws.delete_rows(r)
        _rewrite_voucher_total(ws)
        _del_empty_recon_sheet(wb, sn, ws)
