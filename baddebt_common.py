# -*- coding: utf-8 -*-
"""通用资产科目减值/坏账/跌价准备明细表（2026-08-03 备抵机制通用化）。

任何资产类主科目出现对应准备类科目（坏账准备/跌价准备/减值准备）时复用：
主体×期初余额/本期计提(贷方)/本期转回及转销(借方)/期末余额 + 减值损失科目核对。
供 current_account_detail / longterm_assets_detail / inventory_detail 等 builder 调用。
"""
from collections import defaultdict
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter

_HFILL = PatternFill('solid', fgColor='DDEBF7')
_TFILL = PatternFill('solid', fgColor='D9E1F2')
_BDR = Border(*[Side(style='thin', color='BFBFBF')] * 4)
_F10 = Font(name='Times New Roman', size=10)
_FB = Font(name='Times New Roman', size=10, bold=True)
_FT = Font(name='Times New Roman', size=12, bold=True)
_FR = Font(name='Times New Roman', size=10, bold=True, color='C00000')
_FI = Font(name='Times New Roman', size=9, italic=True, color='808080')
_NUM = '#,##0.00'


def build_baddebt_detail_sheet(wb, tb_full, gl, year, baddebt_name, main_label,
                               loss_kws=('信用减值损失', '资产减值损失'),
                               sheet_name=None):
    """生成『{main_label}减值(坏账/跌价)准备明细表』并返回 True/False（是否有数据）。
    tb_full: read_tb_full 结果 {(e,c,n,y): rec}；gl: read_gl_rows 结果 list。
    baddebt_name: 备抵科目名称关键词（如『坏账准备-应收帐款』『固定资产减值准备』）；
    main_label: 主科目名（用于 sheet 名与标题）；loss_kws: 减值损失科目关键词。"""
    matched = [(e, c, n, v) for (e, c, n, yy), v in (tb_full or {}).items()
               if str(yy) == str(year) and baddebt_name in str(n)]
    if not matched:
        return False
    sn = sheet_name or f'{main_label}减值准备明细表'
    if sn in wb.sheetnames:
        del wb[sn]
    ws = wb.create_sheet(sn)
    ws.cell(1, 1, f'{main_label}减值准备明细表（{year} 年）').font = _FT
    hr = 4
    headers = ['主体', '期初余额', '本期计提', '本期转回及转销', '期末余额']
    for j, h in enumerate(headers, 1):
        cc = ws.cell(hr, j, h)
        cc.font = _FB; cc.fill = _HFILL; cc.border = _BDR
        cc.alignment = Alignment(horizontal='center')
    by_ent = defaultdict(list)
    for e, c, n, v in matched:
        by_ent[e].append(v)
    r = hr + 1
    tot = {'qc': 0.0, 'inc': 0.0, 'dec': 0.0, 'qm': 0.0}
    for ent in sorted(by_ent):
        vs = by_ent[ent]
        # 2026-08-03 符号修复（同往来坏账明细表）：备抵科目 TB 带符号（贷余 qc/qm 为负、
        # 红字 jf/df 用负号）。期初/期末带符号显示；计提=正贷方+红字借方（备抵增加）、
        # 转回及转销=正借方+红字贷方（备抵减少）；勾稽=期初+转回−计提=期末。
        qc = sum(float(v.get('qc') or 0.0) for v in vs)
        qm = sum(float(v.get('qm') or 0.0) for v in vs)
        inc = sum(max(0.0, float(v.get('df') or 0.0)) + max(0.0, -float(v.get('jf') or 0.0))
                  for v in vs)
        dec = sum(max(0.0, float(v.get('jf') or 0.0)) + max(0.0, -float(v.get('df') or 0.0))
                  for v in vs)
        if abs(qc) < 0.005 and abs(inc) < 0.005 and abs(dec) < 0.005 and abs(qm) < 0.005:
            continue
        ws.cell(r, 1, ent)
        for j, val in enumerate((qc, inc, dec, qm), start=2):
            cell = ws.cell(r, j, round(val, 2))
            cell.number_format = _NUM
            cell.alignment = Alignment(horizontal='right')
        for j in range(1, 6):
            ws.cell(r, j).border = _BDR
        tot['qc'] += qc; tot['inc'] += inc; tot['dec'] += dec; tot['qm'] += qm
        r += 1
    ws.cell(r, 1, '合计').font = _FB
    for j, key in enumerate(('qc', 'inc', 'dec', 'qm'), start=2):
        cell = ws.cell(r, j, round(tot[key], 2))
        cell.number_format = _NUM; cell.font = _FB; cell.fill = _TFILL
    for j in range(1, 6):
        ws.cell(r, j).border = _BDR
        ws.cell(r, j).fill = _TFILL
    r += 2
    # —— 计提数与减值损失科目核对（带符号口径：正=计提、红字=转回）——
    ws.cell(r, 1, f'{main_label}准备计提数与减值损失科目核对'
                  f'（借：{" / ".join(loss_kws)} / 贷：{baddebt_name}）：').font = _FB
    r += 1
    for j, h in enumerate(['主体', '年度', '准备贷方发生额', '减值损失借方', '差异'], 1):
        cc = ws.cell(r, j, h); cc.font = _FB
        cc.border = _BDR; cc.fill = _HFILL; cc.alignment = Alignment(horizontal='center')
    r += 1
    prov_by = defaultdict(float)
    loss_by = defaultdict(float)
    for row in (gl or []):
        nm = str(row.get('name') or '')
        if str(row.get('y') or '') != str(year):
            continue
        if baddebt_name in nm:
            prov_by[row.get('e')] += float(row.get('cr') or 0.0)
        if any(k in nm for k in loss_kws):
            loss_by[row.get('e')] += float(row.get('dr') or 0.0)
    for ent in sorted(set(list(prov_by) + list(loss_by))):
        pv = prov_by.get(ent, 0.0); lv = loss_by.get(ent, 0.0)
        ws.cell(r, 1, ent).border = _BDR
        ws.cell(r, 2, year).border = _BDR
        for j, val in enumerate((pv, lv, pv - lv), start=3):
            cell = ws.cell(r, j, round(val, 2))
            cell.number_format = _NUM
            cell.alignment = Alignment(horizontal='right')
            cell.border = _BDR
        if abs(pv - lv) > 0.005:
            ws.cell(r, 5).font = _FR
        r += 1
    ws.cell(r, 1, '注：差异=减值损失借方中含其他构成（审计调整、其他科目减值等），需人工核实。').font = _FI
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for i, w in enumerate([18, 16, 16, 18, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return True


def scan_baddebt_subjects(tb_full, year):
    """扫描 TB 全部备抵科目（名称含 坏账准备/跌价准备/减值准备）。
    返回 [(code, name, 期初绝对值合计, 期末绝对值合计)]（有余额者，跨主体合并）。
    供各 builder 发现『出现该备抵明细 → 补审定表备抵 + 减值准备明细表』的通用机制使用。"""
    out = []
    for (e, c, n, yy), v in (tb_full or {}).items():
        if str(yy) != str(year):
            continue
        nm = str(n)
        if not any(k in nm for k in ('坏账准备', '跌价准备', '减值准备')):
            continue
        qc = float(v.get('qc') or 0.0); qm = float(v.get('qm') or 0.0)
        if abs(qc) < 0.005 and abs(qm) < 0.005:
            continue
        out.append((c, nm, abs(qc), abs(qm)))
    seen = {}
    for c, nm, qc, qm in out:
        if c in seen:
            seen[c][2] += qc; seen[c][3] += qm
        else:
            seen[c] = [c, nm, qc, qm]
    return sorted((tuple(x) for x in seen.values()), key=lambda x: -abs(x[3]))
