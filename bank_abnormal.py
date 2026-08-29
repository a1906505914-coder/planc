# -*- coding: utf-8 -*-
"""银行存款发生额异常分析（剔除银行间）— 审计底稿第2层(科目间关系/发生额分析)。
数据来源：各主体《综合查询明细表》(GL) + 《科目余额表》(TB, 控制数)。
GL 结构要点：银行存款账号多出现在「对方科目」列（如 借财务费用-利息/贷银行存款-工行），
  且「财务费用-利息收入-银行存款活期利息」等损益科目名含"银行存款"字样（非银行账户），须排除。
方法：
  1) 抽取所有触及银行账户的分录行（银行账户在 科目名称 或 对方科目），按凭证聚合去镜像；
  2) 一笔凭证若借贷双方均为本公司银行账户（对手方均为银行）→ 银行间互转，剔除；
  3) 剔除后「外部银行发生额」按对方科目归集，筛查大额/异常分录（一般会计处理之外）。
输出：银行存款发生额异常分析_生成.xlsx
  · 银行存款发生额汇总(剔除银行间)
  · 外部发生额_按对方科目
  · 大额及异常分录
  · 说明
注意：2026 GL 仅 1-5 月(YTD)；TB 为全年，GL缺口=取数范围限制(非差错)。
"""
import paths as P
import os, sys, re
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import audit_common as A

DATA_DIR = P.G
OUT = os.path.join(P.G, '银行存款发生额异常分析_生成.xlsx')
THRESH = 1_000_000.0

THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HFILL = PatternFill('solid', fgColor='D9E1F2')
WARNFILL = PatternFill('solid', fgColor='FFF2CC')
BADFILL = PatternFill('solid', fgColor='FCE4D6')
HFONT = Font(bold=True, name='Times New Roman', size=10)
BFONT = Font(bold=True, name='Times New Roman', size=10)
MONEY = '#,##0.00'

NORMAL_KW = ('应收账款', '应付账款', '预收账款', '预付账款', '库存商品', '原材料', '周转材料',
             '委托加工物资', '在途物资', '应付职工薪酬', '应交税费', '财务费用', '管理费用',
             '销售费用', '制造费用', '研发费用', '主营业务成本', '其他业务成本', '其他应收款',
             '其他应付款', '合同负债', '固定资产', '在建工程', '无形资产', '长期待摊费用',
             '使用权资产', '库存现金', '银行存款', '信用减值损失', '资产减值损失', '营业外支出',
             '营业外收入', '手续费', '利息', '社保', '公积金', '工资', '税金', '成本', '费用')
ABNORMAL_KW = ('借款', '投资', '分红', '股利', '备用金', '保证金', '押金', '往来款', '关联',
               '基金', '理财', '股权', '股本', '实收资本', '资本公积', '捐赠', '个人', '报销',
               '代付', '代偿', '资金拆', '往来')


def _dg(code):
    return re.sub(r'\D', '', str(code))


def is_bank_account(s):
    """严格判定是否为银行账户（排除 财务费用-利息收入-银行存款活期利息 等损益科目）。"""
    if not s:
        return False
    s2 = s.strip()
    if '利息' in s2 or '手续费' in s2 or '费用' in s2:
        return False
    if s2.startswith('银行存款'):
        return True
    if re.search(r'(支行|分行|营业部|信用社|农商|邮储|银行)', s2):
        return True
    if re.search(r'^\d{6,}$', s2):
        return True
    return False


def tb_bank(tb):
    out = {}
    for (e, code, name, y), v in tb.items():
        if _dg(code) == '1002':
            d = out.setdefault((e, y), {'jf': 0.0, 'df': 0.0})
            d['jf'] += v['jf']; d['df'] += v['df']
    if out:
        return out
    for (e, code, name, y), v in tb.items():
        if _dg(code).startswith('1002'):
            d = out.setdefault((e, y), {'jf': 0.0, 'df': 0.0})
            d['jf'] += v['jf']; d['df'] += v['df']
    return out


def classify(rows):
    vouchers = A.group_gl_vouchers(rows)
    ext = []
    inter = []
    for v in vouchers:
        refd = {}
        for l in v['lines']:
            nm = (l.get('name') or '').strip()
            op = (l.get('opp') or '').strip()
            dr = l.get('dr', 0.0) or 0.0
            cr = l.get('cr', 0.0) or 0.0
            if dr == 0 and cr == 0:
                continue
            amt = dr if dr > 0 else cr
            # 确定本行涉及的银行账户与方向
            if is_bank_account(nm):
                ba = nm; direction = '借' if dr > 0 else '贷'; cp = op
            elif is_bank_account(op):
                ba = op; direction = '贷' if dr > 0 else '借'; cp = nm
            else:
                continue
            key = (ba, direction)
            if key not in refd:
                refd[key] = (cp, amt)
        if not refd:
            continue
        nonbank_cp = any(not is_bank_account(cp) for (cp, _) in refd.values())
        base = {'e': v.get('e'), 'y': v.get('y'), 'date': v.get('date'),
                'vtype': v.get('vtype'), 'vno': v.get('vno')}
        if nonbank_cp:
            for (ba, direction), (cp, amt) in refd.items():
                if not is_bank_account(cp):
                    ext.append({**base, 'bank': ba, 'dir': direction, 'amt': amt, 'opp': cp})
        else:
            for (ba, direction), (cp, amt) in refd.items():
                inter.append({**base, 'bank': ba, 'dir': direction, 'amt': amt, 'opp': cp})
    return ext, inter


def flag_line(rec):
    amt = rec['amt']
    direct = rec['dir']
    opp = rec.get('opp') or ''
    tags = []
    if amt >= THRESH:
        tags.append('大额(>=100万)')
        if abs(amt % 1_000_000) < 0.5:
            tags.append('整数大额')
    if any(k in opp for k in ABNORMAL_KW):
        tags.append('可疑对手方(借款/投资/分红/个人/往来等)')
    opp_normal = any(k in opp for k in NORMAL_KW)
    if not opp_normal and amt >= THRESH * 0.2:
        tags.append('对手方非典型(疑异常)')
    return amt, direct, opp, tags


def build(data_dir, out_path):
    entities = A.discover_entities(data_dir)
    if not entities:
        print('❌ 未发现任何账套主体'); return None
    rows = A.read_gl_rows(data_dir, entities)
    tb = A.read_tb_full(data_dir, entities)
    bank_tb = tb_bank(tb)
    years = sorted({y for e, yd in entities.items() for y in yd})
    print(f'· 主体 {len(entities)} 个，年份 {years}，GL 行数 {len(rows)}')

    ext, inter = classify(rows)
    print(f'· 外部银行分录 {len(ext)} 笔，银行间互转 {len(inter)} 笔（已剔除）')

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def style_header(ws, hdr, row=3):
        for c, h in enumerate(hdr, 1):
            cell = ws.cell(row, c, h); cell.font = HFONT; cell.fill = HFILL; cell.border = BORDER
            cell.alignment = Alignment(horizontal='center', wrap_text=True, vertical='center')

    # ---------- Sheet1 汇总 ----------
    ws = wb.create_sheet('银行存款发生额汇总(剔除银行间)')
    ws.cell(1, 1, '银行存款发生额汇总（剔除银行间互转）').font = Font(bold=True, name='Times New Roman', size=10)
    ws.cell(2, 1, '控制数=TB(1002) 借/贷发；GL银行=GL中银行存款借贷方合计；'
                  '外部=GL银行−银行间互转；GL缺口=TB−GL银行(取数范围限制，非差错)。').font = \
        Font(italic=True, name='Times New Roman', size=10, color='808080')
    hdr = ['核算主体', '年度', 'TB借发', 'TB贷发', 'GL银行借发', 'GL银行贷发',
           '银行间借发(剔除)', '银行间贷发(剔除)', '外部借发', '外部贷发', 'GL缺口(借)', '外部占GL借%']
    style_header(ws, hdr, row=3)
    agg = {}
    for r in ext:
        k = (r['e'], r['y'])
        d = agg.setdefault(k, {'ex_d': 0.0, 'ex_c': 0.0})
        if r['dir'] == '借':
            d['ex_d'] += r['amt']
        else:
            d['ex_c'] += r['amt']
    for r in inter:
        k = (r['e'], r['y'])
        d = agg.setdefault(k, {'ex_d': 0.0, 'ex_c': 0.0})
        d.setdefault('it_d', 0.0); d.setdefault('it_c', 0.0)
        if r['dir'] == '借':
            d['it_d'] += r['amt']
        else:
            d['it_c'] += r['amt']
    r = 4
    for e in sorted({k[0] for k in agg}):
        for y in sorted(years):
            k = (e, y)
            d = agg.get(k, {'ex_d': 0.0, 'ex_c': 0.0, 'it_d': 0.0, 'it_c': 0.0})
            bt = bank_tb.get(k, {'jf': 0.0, 'df': 0.0})
            gl_d = d['ex_d'] + d.get('it_d', 0.0); gl_c = d['ex_c'] + d.get('it_c', 0.0)
            pct = (d['ex_d'] / gl_d * 100.0) if gl_d > 1 else 0.0
            row = [e, y, bt['jf'], bt['df'], gl_d, gl_c, d.get('it_d', 0.0), d.get('it_c', 0.0),
                   d['ex_d'], d['ex_c'], bt['jf'] - gl_d, round(pct, 2)]
            for c, val in enumerate(row, 1):
                cell = ws.cell(r, c, val); cell.border = BORDER
                if 3 <= c <= 11:
                    cell.number_format = MONEY
            if abs(bt['jf'] - gl_d) > 1.0:
                ws.cell(r, 11).fill = WARNFILL
            r += 1
    g_btjf = sum(bank_tb.get((e, y), {}).get('jf', 0.0) for e in {k[0] for k in agg} for y in years)
    g_btdf = sum(bank_tb.get((e, y), {}).get('df', 0.0) for e in {k[0] for k in agg} for y in years)
    g_exd = sum(d['ex_d'] for d in agg.values()); g_exc = sum(d['ex_c'] for d in agg.values())
    g_itd = sum(d.get('it_d', 0.0) for d in agg.values()); g_itc = sum(d.get('it_c', 0.0) for d in agg.values())
    g_gld = g_exd + g_itd; g_glc = g_exc + g_itc
    grow = ['全集团合计', '', g_btjf, g_btdf, g_gld, g_glc, g_itd, g_itc, g_exd, g_exc,
            g_btjf - g_gld, round(g_exd / g_gld * 100.0, 2) if g_gld > 1 else 0.0]
    for c, val in enumerate(grow, 1):
        cell = ws.cell(r, c, val); cell.font = BFONT; cell.border = BORDER
        if 3 <= c <= 11:
            cell.number_format = MONEY
    for i, w in enumerate([16, 7, 16, 16, 14, 14, 15, 15, 14, 14, 14, 13], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'A4'

    # ---------- Sheet2 外部按对方科目 ----------
    ws2 = wb.create_sheet('外部发生额_按对方科目')
    ws2.cell(1, 1, '外部银行发生额 — 按对方科目归集（已剔除银行间互转）').font = Font(bold=True, name='Times New Roman', size=10)
    hdr2 = ['核算主体', '年度', '对方科目', '借方笔数', '借方金额', '贷方笔数', '贷方金额', '净额', '标记']
    style_header(ws2, hdr2, row=3)
    cp = {}
    for r0 in ext:
        key = (r0['e'], r0['y'], r0['opp'] or '（空）')
        d = cp.setdefault(key, {'nd': 0, 'sd': 0.0, 'nc': 0, 'sc': 0.0})
        if r0['dir'] == '借':
            d['nd'] += 1; d['sd'] += r0['amt']
        else:
            d['nc'] += 1; d['sc'] += r0['amt']
    r = 4
    for key in sorted(cp.keys(), key=lambda x: -(cp[x]['sd'] + cp[x]['sc'])):
        e, y, opp = key; d = cp[key]
        net = d['sd'] - d['sc']
        tags = []
        if d['sd'] + d['sc'] >= THRESH * 5:
            tags.append('大额对手方(>=500万)')
        if any(k in opp for k in ABNORMAL_KW):
            tags.append('可疑对手方')
        opp_normal = any(k in opp for k in NORMAL_KW)
        if not opp_normal and (d['sd'] + d['sc']) >= THRESH:
            tags.append('对手方非典型')
        row = [e, y, opp, d['nd'], d['sd'], d['nc'], d['sc'], net, '；'.join(tags)]
        for c, val in enumerate(row, 1):
            cell = ws2.cell(r, c, val); cell.border = BORDER
            if c in (5, 7, 8):
                cell.number_format = MONEY
            if c == 9 and tags:
                cell.fill = BADFILL if ('可疑' in ''.join(tags) or '非典型' in ''.join(tags)) else WARNFILL
        r += 1
    for i, w in enumerate([14, 7, 34, 10, 16, 10, 16, 16, 24], 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws2.freeze_panes = 'A4'

    # ---------- Sheet3 大额及异常分录 ----------
    ws3 = wb.create_sheet('大额及异常分录')
    ws3.cell(1, 1, '大额及异常银行分录清单（单笔 >= 100万 或 对手方可疑/非典型）').font = Font(bold=True, name='Times New Roman', size=10)
    hdr3 = ['核算主体', '年度', '日期', '凭证字', '凭证号', '银行账户', '方向', '金额', '对方科目', '标记']
    style_header(ws3, hdr3, row=3)
    flagged = []
    for r0 in ext:
        amt, direct, opp, tags = flag_line(r0)
        if tags:
            flagged.append((amt, r0, direct, opp, tags))
    flagged.sort(key=lambda x: -abs(x[0]))
    r = 4
    for amt, r0, direct, opp, tags in flagged:
        row = [r0['e'], r0['y'], r0['date'], r0['vtype'], r0['vno'], r0['bank'], direct, amt, opp, '；'.join(tags)]
        for c, val in enumerate(row, 1):
            cell = ws3.cell(r, c, val); cell.border = BORDER
            if c == 8:
                cell.number_format = MONEY
            if c == 10:
                cell.fill = BADFILL if ('可疑' in ''.join(tags) or '非典型' in ''.join(tags)) else WARNFILL
        r += 1
    if not flagged:
        ws3.cell(r, 1, '无达到阈值或可疑对手方的大额银行分录。').font = Font(name='Times New Roman', italic=True)
    for i, w in enumerate([14, 7, 12, 8, 10, 30, 6, 16, 30, 30], 1):
        ws3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws3.freeze_panes = 'A4'

    # ---------- Sheet4 说明 ----------
    ws4 = wb.create_sheet('说明')
    notes = [
        '银行存款发生额异常分析（剔除银行间） 编制与勾稽说明',
        '',
        '一、方法（第2层：科目间发生额分析）',
        '  1. 抽取《综合查询明细表》所有触及银行账户的分录行（银行账户在科目名称或对方科目列），按凭证聚合并去除借贷镜像；',
        '     （注意：GL 中「财务费用-利息收入-银行存款活期利息」等损益科目名含"银行存款"字样，已排除，不视为银行账户。）',
        '  2. 一笔凭证借贷双方均为本公司银行账户（对手方均为银行）→ 判定为银行间互转，剔除；',
        '  3. 剔除后「外部银行发生额」按对方科目归集，筛查大额/异常分录（一般会计处理之外）。',
        '',
        '二、与 TB 控制数勾稽',
        '  · TB(1002) 借/贷发 为权威控制数；GL银行 = GL 中银行存款借贷方合计；',
        '  · GL缺口 = TB − GL银行：因综合查询明细表为非全量抽取，GL银行 ≤ TB，缺口=取数范围限制，非账务差错。',
        '',
        '三、异常判定规则（阈值 100万元）',
        '  · 大额：单笔 >= 100万（整数大额另行标注）；',
        '  · 可疑对手方：对方科目含 借款/投资/分红/股利/备用金/保证金/往来款/关联/理财/股权/个人 等；',
        '  · 对手方非典型：对方科目不在常规银行收支对手方清单(应收/应付/存货/薪酬/税费/费用/现金等)且金额>=20万。',
        '  上述清单为筛选提示，需结合凭证与业务实质进一步核查是否构成会计差错。',
        '',
        '四、口径提示',
        '  · 2026 综合查询明细表仅含 1-5 月(YTD)，外部发生额仅反映年初至5月；TB 为全年，缺口偏大属正常。',
        f'  · 银行间互转共识别 {len(inter)} 笔已剔除；外部银行分录 {len(ext)} 笔进入分析。',
    ]
    for i, t in enumerate(notes, 1):
        ws4.cell(i, 1, t)
        if i == 1:
            ws4.cell(i, 1).font = Font(name='Times New Roman', bold=True, size=10)
    ws4.column_dimensions['A'].width = 110

    path, warn = A._safe_save(wb, out_path)
    if warn:
        print('  ⚠️', warn)
    print(f'  ✓ 已保存：{path}  (外部借发={g_exd:,.2f}, 外部贷发={g_exc:,.2f}, 银行间剔除={len(inter)}笔)')
    return path


def main():
    data_dir = sys.argv[1] if len(sys.argv) > 1 else DATA_DIR
    out = sys.argv[2] if len(sys.argv) > 2 else OUT
    build(data_dir, out)


if __name__ == '__main__':
    main()
