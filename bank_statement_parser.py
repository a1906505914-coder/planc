# -*- coding: utf-8 -*-
"""bank_statement_parser.py —— 银行网银明细解析器（注册表模式）

原则（用户 2026-08-13）：网银解析必须能适应所有账套/所有银行 → 每家银行一条解析规则
注册在 BANK_RULES，新银行只需加一个解析函数，主流程零改动。

统一输出行结构（企业视角，与 read_gl_bank 方向对齐）：
    {'bank': 银行名, 'file': 源文件名, 'account_no': 账号,
     'date': 'YYYY-MM-DD', 'direction': '收'/'支',
     'amount': float, 'counterparty': 对方户名, 'summary': 摘要, 'balance': float}

方向语义（银行角度 → 企业角度）：
    银行「贷方/收入/转入/来账」 = 企业账「借方增加」 → direction='收'
    银行「借方/支出/转出/往账」 = 企业账「贷方减少」 → direction='支'

已实测 9 个银行文件（6100，2026-08-13）：
  中信 xlsx  表头行15；账号 行5B
  中行 xls   表头行8；账号 行1B；col0 来账=收/往账=支；日期 YYYYMMDD
  建行 xls   表头行0；账号 col0
  交行 xls   表头行1；账号 行0B
  工行 xlsx  表头行1；借贷标志 col3（贷=收/借=支）；转出 col8/转入 col9
  农行 xls   表头行2；账号 行1B；收入 col1/支出 col2
  国开行 xls 表头行0；账号 col0；借方(支取) col3/贷方(收入) col4
"""
import paths as P
import os
import re
import datetime

# ---------------------------------------------------------------------------
# 通用清洗
# ---------------------------------------------------------------------------
def _clean_amt(v):
    """'50,000.00' / 50000.0 / '0.04' → float（None/'' → 0.0）"""
    if v is None:
        return 0.0
    s = str(v).strip().replace(',', '').replace('，', '').replace(' ', '')
    if not s or s in ('-', '--', 'None'):
        return 0.0
    try:
        return float(s)
    except ValueError:
        m = re.search(r'-?\d+(\.\d+)?', s)
        return float(m.group()) if m else 0.0


def _norm_date(v):
    """'2026-07-13' / '20260108' / '20260108 05:13' / '2026-01-21 15:' → '2026-01-21'；
    无法解析（合计/页脚行等）→ ''（调用方过滤）。"""
    if v is None:
        return ''
    s = str(v).strip()
    m = re.match(r'^(\d{4})[-/年.]?(\d{1,2})[-/月.]?(\d{1,2})', s)
    if m:
        return '%s-%02d-%02d' % (m.group(1), int(m.group(2)), int(m.group(3)))
    m = re.match(r'^(\d{4})(\d{2})(\d{2})', s)
    if m:
        return '%s-%s-%s' % (m.group(1), m.group(2), m.group(3))
    return ''


def _rows_of_xls(path):
    import xlrd
    wb = xlrd.open_workbook(path)
    ws = wb.sheet_by_index(0)
    return [[ws.cell_value(i, j) for j in range(ws.ncols)] for i in range(ws.nrows)]


def _rows_of_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_of_csv(path):
    """csv → 行列表（编码自适应：工行 GBK / 中行 UTF-8 BOM）。
    ⚡ 2026-08-13：用标准 csv.reader（自动处理引号内逗号），勿按制表符切。
    ⚡⚡ 2026-08-14：OLE 头探测——某些文件扩展名 .CSV 实为老式 .xls 复合文档
        （3800 工银实证：d0cf11e0 = OLE），按 xls 读取而非 csv（否则二进制乱码）。"""
    with open(path, 'rb') as _f:
        _h = _f.read(8)
    if _h[:4] == b'\xd0\xcf\x11\xe0':  # OLE 复合文档（老式 xls）
        return _rows_of_xls(path)
    import csv as _csv
    last_err = None
    for enc in ('utf-8-sig', 'gbk'):
        try:
            with open(path, 'r', encoding=enc) as f:
                return [[c.strip() for c in row] for row in _csv.reader(f)]
        except (UnicodeDecodeError, UnicodeError) as e:
            last_err = e
            continue
    with open(path, 'r', encoding='gbk', errors='replace') as f:
        return [[c.strip() for c in row] for row in _csv.reader(f)]


# ---------------------------------------------------------------------------
# 各银行解析器（返回统一行列表）
# ---------------------------------------------------------------------------
def parse_zhongxin(path, rows):
    """中信：动态探测表头行 + 按列名映射（兼容新旧两格式，2026-08-15）——
    ① 旧格式（表头行16）：交易日期|交易时间|对方账号|对方账户名称|…|借方发生额|贷方发生额|账户余额|摘要|…|附言
       借方 col5 / 贷方 col6 / 余额 col7 / 对方名 col3 / 摘要 col8+col12；
    ② 新变体（3100 中信2901，表头行13）：账号|账号名称|币种|交易日|…|借方金额|贷方金额|余额|摘要|…
       借方 col7 / 贷方 col8 / 余额 col9 / 交易日 col3 / 摘要 col10。
    原固定表头行15+固定列 → 2901 表头行13/列错位 → 0 笔（累计贷 2 笔 0.06×2 结息全漏）。
    """
    hi = None
    for i, r in enumerate(rows[:20]):
        if not r:
            continue
        s = '|'.join(str(x) if x is not None else '' for x in r[:15])
        if '借方' in s and '贷方' in s:
            hi = i
            break
    if hi is None:
        return []
    hdr = rows[hi]
    _h = [str(x) if x is not None else '' for x in hdr[:15]]

    def _ci(*kws):
        for j, hh in enumerate(_h):
            if any(k in hh for k in kws):
                return j
        return None

    c_date = _ci('交易日', '交易日期')
    c_dj = _ci('借方')
    c_cr = _ci('贷方')
    c_bal = _ci('余额')
    c_sm = _ci('摘要')
    c_cp = _ci('对方账户名称', '对方名称', '户名')
    if c_date is None or (c_dj is None and c_cr is None):
        return []
    acct = str(rows[5][1]).strip() if len(rows) > 5 else ''
    out = []
    for r in rows[hi + 1:]:
        if not r or len(r) <= c_date or not str(r[c_date]).strip():
            continue
        dj = _clean_amt(r[c_dj]) if c_dj is not None and len(r) > c_dj else 0.0
        cr = _clean_amt(r[c_cr]) if c_cr is not None and len(r) > c_cr else 0.0
        if dj == 0 and cr == 0:
            continue
        date = _norm_date(str(r[c_date])[:19])
        if not date:
            continue
        sm = str(r[c_sm]).strip() if c_sm is not None and len(r) > c_sm else ''
        if c_sm is not None and len(r) > c_sm + 4:      # 旧格式附言在 col12（=摘要+4）
            sm = (sm + ' ' + str(r[c_sm + 4]).strip()).strip()
        out.append({
            'bank': '中信', 'file': os.path.basename(path), 'account_no': acct,
            'date': date, 'direction': '支' if dj else '收',
            'amount': dj if dj else cr,
            'counterparty': str(r[c_cp]).strip() if c_cp is not None and len(r) > c_cp else '',
            'summary': sm[:60],
            'balance': _clean_amt(r[c_bal]) if c_bal is not None and len(r) > c_bal else None,
        })
    return out


def parse_boc(path, rows):
    """中行：表头行8；账号 行1B；col0 来账=收/往账=支；日期 col10(YYYYMMDD)；金额 col13；余额 col14"""
    out = []
    acct = str(rows[1][1]).strip() if len(rows) > 1 else ''
    for r in rows[9:]:
        if not r or not str(r[0]).strip():
            continue
        typ = str(r[0]).strip()
        amt = abs(_clean_amt(r[13])) if len(r) > 13 else 0.0  # ⚡ 中行往账金额带负号，取绝对值
        if amt == 0:
            continue
        direction = '收' if '来' in typ or '收' in typ else ('支' if '往' in typ or '付' in typ else '')
        if not direction:
            continue
        cp = str(r[9]).strip() if len(r) > 9 else ''
        out.append({
            'bank': '中行', 'file': os.path.basename(path), 'account_no': acct,
            'date': _norm_date(r[10]), 'direction': direction, 'amount': amt,
            'counterparty': cp, 'summary': str(r[23]).strip() if len(r) > 23 else '',
            'balance': _clean_amt(r[14]) if len(r) > 14 else None,
        })
    return out


def parse_ccb(path, rows):
    """建行：表头行0；账号 col0；借方(支取) col3 / 贷方(收入) col4；余额 col5；对方 col7；摘要 col11"""
    out = []
    for r in rows[1:]:
        if not r or len(r) < 5 or not str(r[0]).strip():
            continue
        dj = _clean_amt(r[3])
        cr = _clean_amt(r[4])
        if dj == 0 and cr == 0:
            continue
        out.append({
            'bank': '建行', 'file': os.path.basename(path), 'account_no': str(r[0]).strip(),
            'date': _norm_date(r[2]), 'direction': '支' if dj else '收',
            'amount': dj if dj else cr,
            'counterparty': str(r[7]).strip() if len(r) > 7 else '',
            'summary': str(r[11]).strip() if len(r) > 11 else '',
            'balance': _clean_amt(r[5]) if len(r) > 5 else None,
        })
    return out


def parse_bocom(path, rows):
    """交行：表头行1；账号 行0B；借方(支出) col1 / 贷方(收入) col2；余额 col3；对方名 col5；摘要 col6"""
    out = []
    acct = str(rows[0][1]).strip() if rows else ''
    for r in rows[2:]:
        if not r or not str(r[0]).strip():
            continue
        dj = _clean_amt(r[1]) if len(r) > 1 else 0.0
        cr = _clean_amt(r[2]) if len(r) > 2 else 0.0
        if dj == 0 and cr == 0:
            continue
        out.append({
            'bank': '交行', 'file': os.path.basename(path), 'account_no': acct,
            'date': _norm_date(r[0]), 'direction': '支' if dj else '收',
            'amount': dj if dj else cr,
            'counterparty': str(r[5]).strip() if len(r) > 5 else '',
            'summary': str(r[6]).strip() if len(r) > 6 else '',
            'balance': _clean_amt(r[3]) if len(r) > 3 else None,
        })
    return out


def parse_icbc_asia(path, rows):
    """工银亚洲（中国香港）：繁体表头，多账户多币种。
    列：0账户名 / 3账号 / 5业务类型 / 6交易时间 / 8扣账/入账 / 9币种 / 10金额 /
        11余额 / 12摘要 / 13对方账号 / 14对方户名 / 15凭证号。
    方向：入账=收 / 扣账=支。返回行带 account_no（col3）+ currency（col9）。
    ⚡ 2026-08-14 3800 实证：LANTEAN HOLDING 港元/人民币/美元多账户，工银亚洲。
    表头行 0（含繁体列名）；数据行 1 起。"""
    out = []
    acct_default = ''
    for r in rows[1:]:
        if not r or len(r) < 11:
            continue
        biz = str(r[5] or '').strip() if len(r) > 5 else ''
        tt = str(r[6] or '').strip() if len(r) > 6 else ''
        dv = str(r[8] or '').strip() if len(r) > 8 else ''
        cur = str(r[9] or '').strip() if len(r) > 9 else ''
        amt = _clean_amt(r[10]) if len(r) > 10 else 0.0
        if amt == 0:
            continue
        date = _norm_date(tt.split(' ')[0]) if tt else ''
        acct = str(r[3] or '').strip() if len(r) > 3 else acct_default
        direction = '收' if '入' in dv else ('支' if '扣' in dv else '')
        if not direction:
            continue
        cp = str(r[14] or '').strip() if len(r) > 14 else ''
        summ = str(r[12] or '').strip() if len(r) > 12 else ''
        out.append({
            'bank': '工行', 'file': os.path.basename(path),
            'account_no': acct, 'currency': cur,
            'date': date, 'direction': direction, 'amount': amt,
            'counterparty': cp, 'summary': summ,
            'balance': _clean_amt(r[11]) if len(r) > 11 else None,
            'biz': biz,
        })
    return out


def parse_icbc(path, rows):
    """工行：表头行1；借贷标志 col3（贷=收/借=支）；对方 col4；转出 col8 / 转入 col9；账号从文件名（N账号）"""
    out = []
    m = re.search(r'（(\d+)账号）', os.path.basename(path))
    acct = m.group(1) if m else ''
    for r in rows[2:]:
        if not r or len(r) < 10 or not str(r[2]).strip():
            continue
        flag = str(r[3]).strip()
        zc = _clean_amt(r[8])
        zr = _clean_amt(r[9])
        if zc == 0 and zr == 0:
            continue
        direction = '收' if '贷' in flag or zr else '支'
        out.append({
            'bank': '工行', 'file': os.path.basename(path), 'account_no': acct,
            'date': _norm_date(r[2]), 'direction': direction,
            'amount': zr if '贷' in flag else zc,
            'counterparty': str(r[4]).strip() if len(r) > 4 else '',
            'summary': (str(r[6]).strip() + ' ' + str(r[7]).strip()).strip() if len(r) > 7 else '',
            'balance': None,
        })
    return out


def parse_abc(path, rows):
    """农行：表头行2；账号 行1B('账号:10-5456010400')；收入 col1 / 支出 col2；余额 col3；对方名 col5；摘要 col7"""
    out = []
    acct = ''
    if len(rows) > 1:
        m = re.search(r'账号[:：]?\s*(\S+)', str(rows[1][0]))
        acct = m.group(1) if m else ''
    for r in rows[3:]:
        if not r or not str(r[0]).strip():
            continue
        # ⚡ 2026-08-13 修复：过滤合计/页脚行（'5 收 91,007,497.02' 首列非日期）
        if not _norm_date(r[0]):
            continue
        inc = _clean_amt(r[1]) if len(r) > 1 else 0.0
        exp = _clean_amt(r[2]) if len(r) > 2 else 0.0
        if inc == 0 and exp == 0:
            continue
        out.append({
            'bank': '农行', 'file': os.path.basename(path), 'account_no': acct,
            'date': _norm_date(r[0]), 'direction': '收' if inc else '支',
            'amount': inc if inc else exp,
            'counterparty': str(r[5]).strip() if len(r) > 5 else '',
            'summary': str(r[7]).strip() if len(r) > 7 else '',
            'balance': _clean_amt(r[3]) if len(r) > 3 else None,
        })
    return out


def parse_cdb(path, rows):
    """国开行：表头行0；账号 col0；借方(支取) col3 / 贷方(收入) col4；余额 col5；对方 col7"""
    out = []
    for r in rows[1:]:
        if not r or len(r) < 5 or not str(r[0]).strip():
            continue
        dj = _clean_amt(r[3])
        cr = _clean_amt(r[4])
        if dj == 0 and cr == 0:
            continue
        out.append({
            'bank': '国开行', 'file': os.path.basename(path), 'account_no': str(r[0]).strip(),
            'date': _norm_date(r[2]), 'direction': '支' if dj else '收',
            'amount': dj if dj else cr,
            'counterparty': str(r[7]).strip() if len(r) > 7 else '',
            'summary': str(r[10]).strip() if len(r) > 10 else '',
            'balance': _clean_amt(r[5]) if len(r) > 5 else None,
        })
    return out


def parse_nsbank(path, rows):
    """农商行：表头行4（交易日期/对方户名/收入金额/支出金额/余额/账号/币种/对方账号/用途/附言/核心流水号）；
    账号 行1 col1；收入 col2 / 支出 col3；余额 col4；对方户名 col1；用途 col8 + 附言 col9"""
    out = []
    acct = ''
    if len(rows) > 1:
        m = re.search(r'(\d{10,})', str(rows[1][1])) if len(rows[1]) > 1 else None
        acct = m.group(1) if m else ''
    for r in rows[5:]:
        if not r or len(r) < 5 or not str(r[0]).strip():
            continue
        inc = _clean_amt(r[2]) if len(r) > 2 else 0.0
        exp = _clean_amt(r[3]) if len(r) > 3 else 0.0
        if inc == 0 and exp == 0:
            continue
        usages = []
        if len(r) > 8 and str(r[8]).strip():
            usages.append(str(r[8]).strip())
        if len(r) > 9 and str(r[9]).strip():
            usages.append(str(r[9]).strip())
        out.append({
            'bank': '农商行', 'file': os.path.basename(path), 'account_no': acct,
            'date': _norm_date(r[0]), 'direction': '收' if inc else '支',
            'amount': inc if inc else exp,
            'counterparty': str(r[1]).strip() if len(r) > 1 else '',
            'summary': ' '.join(usages)[:60],
            'balance': _clean_amt(r[4]) if len(r) > 4 else None,
        })
    return out


def parse_icbc_csv(path, rows):
    """工行 csv（GBK/制表符）：两种格式自动兼容——
    ① 3100 格式：行0 [HISTORYDETAIL]+行1 表头（凭证号/本方账号/对方账号/交易时间/借·贷/
       借方发生额/贷方发生额/…/对方单位名称/余额）；借=支出、贷=收入。
    ② 3000 格式（2026-08-14 新发现）：表头（对方账号/交易时间/借贷标志/对方单位/对方行号/
       用途/摘要/附言/发生额/余额/币种）；借贷标志单列（借=支/贷=收）+ 发生额单列。
    原代码只认 ① 的列索引（r[3]日期/r[4]借贷/r[5][6]金额）→ 3000 格式整列错位解析 0 笔。
    """
    out = []
    hi, hdr = _probe_hdr(rows, ['交易时间'])
    if hi is None:
        return out
    _hdrs = '|'.join(str(x) for x in hdr[:15])
    # ⚡⚡ 2026-08-15 ③ 3500 格式（无借贷标志）：对方账号|交易时间|对方单位|用途|摘要|附言|
    #   回单个性化信息|余额|转出金额|转入金额 —— 金额分 转出/转入 两列（转出=支、转入=收），
    #   无借贷标志列（原走①分支 r[3]='用途'空 → 全跳过 0 笔，3500 工行 11 份月拆/半月拆全漏）。
    if '借贷标志' not in _hdrs and ('转出金额' in _hdrs and '转入金额' in _hdrs):
        dc = _col(hdr, '交易时间')
        inc_c = _col(hdr, '转入金额') or _col(hdr, '收入')
        exp_c = _col(hdr, '转出金额') or _col(hdr, '支出')
        cpc = _col(hdr, '对方单位') or _col(hdr, '对方户名') or _col(hdr, '户名')
        smc = _col(hdr, '摘要')
        balc = _col(hdr, '余额')
        acct_c = _col(hdr, '本方账号') or _col(hdr, '账号')
        if dc is None or (inc_c is None and exp_c is None):
            return out
        for r in rows[hi + 1:]:
            if not r or len(r) <= dc or not str(r[dc]).strip():
                continue
            date = _norm_date(str(r[dc])[:19])
            if not date:
                continue
            zr = _clean_amt(r[inc_c]) if inc_c is not None and len(r) > inc_c else 0.0
            zc = _clean_amt(r[exp_c]) if exp_c is not None and len(r) > exp_c else 0.0
            if zr == 0 and zc == 0:
                continue
            direction = '收' if zr else '支'
            amt = zr if zr else zc
            out.append({
                'bank': '工行', 'file': os.path.basename(path),
                'account_no': str(r[acct_c]).strip('"') if acct_c is not None and len(r) > acct_c else str(r[0]).strip('"'),
                'date': date, 'direction': direction, 'amount': amt,
                'counterparty': str(r[cpc]).strip() if cpc is not None and len(r) > cpc else '',
                'summary': str(r[smc]).strip()[:60] if smc is not None and len(r) > smc else '',
                'balance': _clean_amt(r[balc]) if balc is not None and len(r) > balc else None,
            })
        return out
    if '借贷标志' in _hdrs:
        # ② 3000 格式：借贷标志 + 发生额单列
        # ②b 3400 格式（2026-08-14）：借贷标志 + 转入金额/转出金额两列（贷=转入收/借=转出支）
        dc = _col(hdr, '交易时间'); dirc = _col(hdr, '借贷标志')
        ac = _col(hdr, '发生额')
        inc_c = _col(hdr, '转入金额') or _col(hdr, '收入')
        exp_c = _col(hdr, '转出金额') or _col(hdr, '支出')
        cpc = _col(hdr, '对方单位'); smc = _col(hdr, '摘要'); balc = _col(hdr, '余额')
        if dc is None or dirc is None or (ac is None and inc_c is None and exp_c is None):
            return out
        for r in rows[hi + 1:]:
            if not r or len(r) <= max(dc, dirc) or not str(r[dc]).strip():
                continue
            date = _norm_date(str(r[dc])[:19])
            if not date:
                continue
            fl = str(r[dirc]).strip()
            if '借' in fl:
                direction = '支'
            elif '贷' in fl:
                direction = '收'
            else:
                continue
            if ac is not None:
                amt = _clean_amt(r[ac])
            elif direction == '收':
                amt = _clean_amt(r[inc_c]) if inc_c is not None and len(r) > inc_c else 0
            else:
                amt = _clean_amt(r[exp_c]) if exp_c is not None and len(r) > exp_c else 0
            if amt <= 0:
                continue
            out.append({
                'bank': '工行', 'file': os.path.basename(path), 'account_no': str(r[0]).strip('"'),
                'date': date, 'direction': direction, 'amount': amt,
                'counterparty': str(r[cpc]).strip() if len(r) > cpc else '',
                'summary': str(r[smc]).strip()[:60] if smc is not None and len(r) > smc else '',
                'balance': _clean_amt(r[balc]) if balc is not None and len(r) > balc else None,
            })
        return out
    # ① 3100 格式：借贷标志 + 借方/贷方两列
    for r in rows[2:]:
        if len(r) < 7 or not str(r[3]).strip():
            continue
        date = _norm_date(str(r[3])[:19])
        if not date:
            continue
        dr = '借' if '借' in str(r[4]) else ('贷' if '贷' in str(r[4]) else '')
        if not dr:
            continue
        inc = _clean_amt(r[6])
        exp = _clean_amt(r[5])
        amt = inc if dr == '贷' else exp
        if amt <= 0:
            continue
        out.append({
            'bank': '工行', 'file': os.path.basename(path), 'account_no': str(r[1]).strip('"'),
            'date': date, 'direction': '收' if dr == '贷' else '支', 'amount': amt,
            'counterparty': str(r[10]).strip() if len(r) > 10 else '',
            'summary': str(r[8]).strip()[:60] if len(r) > 8 else '',
            'balance': _clean_amt(r[11]) if len(r) > 11 else None,
        })
    return out


# ⚡ 2026-08-13 配置表驱动通用解析器（3100 网银 15 家新银行一次性注册）：
#   cfg = dict(skip=表头行数, date=日期列, inc=收入列, exp=支出列, cp=对方名列,
#              sm=摘要列, bal=余额列, dir=方向列(优先) 或 'zn'=按本方户名判断)
def _grow(path, rows, cfg, bank):
    out = []
    for r in rows[cfg.get('skip', 1):]:
        if not r:
            continue
        inc = _clean_amt(r[cfg['inc']]) if cfg.get('inc') is not None and len(r) > cfg['inc'] else 0.0
        exp = _clean_amt(r[cfg['exp']]) if cfg.get('exp') is not None and len(r) > cfg['exp'] else 0.0
        if inc == 0 and exp == 0:
            continue
        dc = cfg.get('date', 0)
        date = _norm_date(str(r[dc])[:12]) if len(r) > dc else ''
        if not date:
            continue
        if cfg.get('dir'):
            s = str(r[cfg['dir']]) if len(r) > cfg['dir'] else ''
            direction = '收' if ('收' in s or '贷' in s) else ('支' if ('支' in s or '借' in s) else ('收' if inc else '支'))
        elif cfg.get('zn'):
            me = str(r[cfg['zn']]) if len(r) > cfg['zn'] else ''
            direction = '支' if '国望' in me else ('收' if inc else '支')
        else:
            direction = '收' if inc else '支'
        amt = inc if direction == '收' else exp
        if amt <= 0:
            continue
        # ⚡ 2026-08-15 cp2：收/支对方列不同（平安 CSV 付款人名称↔收款人名称）
        cp_i = cfg.get('cp') if direction == '收' else cfg.get('cp2', cfg.get('cp'))
        out.append({
            'bank': bank, 'file': os.path.basename(path), 'account_no': '',
            'date': date, 'direction': direction, 'amount': amt,
            'counterparty': str(r[cp_i]).strip() if cp_i is not None and len(r) > cp_i else '',
            'summary': str(r[cfg['sm']]).strip()[:60] if cfg.get('sm') is not None and len(r) > cfg['sm'] else '',
            'balance': _clean_amt(r[cfg['bal']]) if cfg.get('bal') is not None and len(r) > cfg['bal'] else None,
        })
    return out


def _probe_hdr(rows, kws):
    """找同时含全部关键词的表头行 → (行号, 表头列表)；找不到 → (None, None)。"""
    for i, r in enumerate(rows):
        joined = '|'.join(str(x) for x in r[:25])
        if all(k in joined for k in kws):
            return i, [str(x) for x in r[:25]]
    return None, None


def _col(hdr, *kws):
    for k in kws:
        for j, h in enumerate(hdr):
            if k in h:
                return j
    return None


def parse_probed(path, rows, bank, hdr_kws, inc_kw, exp_kw, date_kw, cp_kw=None, sm_kw=None):
    """表头探测通用解析：自动定位日期/收入/支出/对方名 列（民生/苏州等表头位置不定的银行）。"""
    hi, hdr = _probe_hdr(rows, hdr_kws)
    if hi is None:
        return []
    dc, ic, ec = _col(hdr, date_kw), _col(hdr, inc_kw), _col(hdr, exp_kw)
    if dc is None or (ic is None and ec is None):
        return []
    return _grow(path, rows, {'skip': hi + 1, 'date': dc, 'inc': ic, 'exp': ec,
                              'cp': _col(hdr, cp_kw) if cp_kw else None,
                              'sm': _col(hdr, sm_kw) if sm_kw else None}, bank)


def parse_shanghai(path, rows):
    """上海银行：表头行5（交易流水号/交易时间/记账日期/交易方向/交易金额/余额/对手账号）"""
    return _grow(path, rows, {'skip': 6, 'date': 3, 'dir': 4, 'inc': 5, 'exp': 5, 'cp': 7, 'bal': 6}, '上海银行')


def parse_gd_eb(path, rows):
    """光大银行：表头行6（交易日期/…/借方金额（支出）/贷方金额（收入）/账户余额/对方账号/对方名称/摘要）"""
    return _grow(path, rows, {'skip': 7, 'date': 0, 'exp': 2, 'inc': 3, 'bal': 4, 'cp': 6, 'sm': 7}, '光大银行')


def parse_cib(path, rows):
    """兴业银行：表头行1；日期从流水编号(col0 20260710xxxx)；借方7/贷方8；户名3；摘要10；对方11"""
    return _grow(path, rows, {'skip': 2, 'date': 0, 'exp': 7, 'inc': 8, 'cp': 3, 'sm': 10, 'bal': 9}, '兴业银行')


def parse_bob(path, rows):
    """北京银行：表头行2（序号/交易时间/币种/借方发生额/贷方发生额/余额/对手方/对手方账号）"""
    return _grow(path, rows, {'skip': 3, 'date': 1, 'exp': 3, 'inc': 4, 'bal': 5, 'cp': 6}, '北京银行')


def parse_huaxia(path, rows):
    """华夏银行：表头（序号/交易日期/交易时间/支出金额/收入金额/余额/对方账号/对方户名）
    ⚡ 2026-08-14 修复：表头行号不固定——1-6 月文件表头行9、7 月文件多两行汇总（收入总
    金额/支出总金额）表头行11；原写死 skip=10 把 7 月『支出总金额 14928.64』当数据行 →
    日期解析成 1492-08-64、金额 1.00 假差异。改表头探测。"""
    return parse_probed(path, rows, '华夏银行', ['交易日期', '序号'],
                        '收入金额', '支出金额', '交易日期', cp_kw='对方户名')


def parse_njbank(path, rows):
    """南京银行：表头行2（交易日期/收入/支出/账户余额/对方账号/对方户名/对方行名/摘要）"""
    return _grow(path, rows, {'skip': 3, 'date': 0, 'inc': 1, 'exp': 2, 'bal': 3, 'cp': 5, 'sm': 7}, '南京银行')


def parse_nbbank(path, rows):
    """宁波银行：表头行3（流水号/交易账号/交易日期/…/借方金额/贷方金额）"""
    return _grow(path, rows, {'skip': 4, 'date': 2, 'exp': 6, 'inc': 7}, '宁波银行')


def parse_pingan(path, rows):
    """平安银行：两格式自动探测。
    ① 新版 CSV（3200 实证）：表头(交易日期/传票号/借方/贷方/余额/付款人账户/付款人名称/收款人账户/收款人名称/用途/摘要)
       ——借方=支/贷方=收，收方对方=付款人名称、支方对方=收款人名称（cp2）。
       ⚡ 2026-08-15 修复：旧配置 inc=3/exp=4/bal=6 把『余额』列当支出列 → 大额支出被
       误取成余额数（3200 实证：支出 50,000,000 被解析成 187,311.57）。
    ② 旧版 xls：表头行0（交易时间/账号/币种/收入/支出/冲正标志/账户余额/对方账号）。"""
    hi, hdr = _probe_hdr(rows, ['借方', '贷方', '余额'])
    if hi is not None:
        return _grow(path, rows, {'skip': hi + 1, 'date': _col(hdr, '交易日期'),
                                  'inc': _col(hdr, '贷方'), 'exp': _col(hdr, '借方'),
                                  'bal': _col(hdr, '余额'),
                                  'cp': _col(hdr, '付款人名称'), 'cp2': _col(hdr, '收款人名称'),
                                  'sm': _col(hdr, '摘要')}, '平安银行')
    return _grow(path, rows, {'skip': 1, 'date': 0, 'inc': 3, 'exp': 4, 'bal': 6, 'cp': 7}, '平安银行')


def parse_cmb(path, rows):
    """招商银行：表头探测（账号/账号名称/币种/交易日/交易时间/起息日/交易类型/借方金额/贷方金额/…）
    ⚡ 2026-08-14 修复：表头行号不固定——3100=行12、3000=行5 → 原写死 skip=13 导致 3000
    招行（15 行表）跳过大部分数据只解析出 2/4 笔。"""
    return parse_probed(path, rows, '招商银行', ['交易日', '借方金额'],
                        '贷方金额', '借方金额', '交易日')


def parse_psbc(path, rows):
    """邮储银行：表头行5（序号/交易日期/…/支出金额/收入金额/余额/对方账号）"""
    return _grow(path, rows, {'skip': 6, 'date': 1, 'exp': 4, 'inc': 5, 'bal': 6, 'cp': 7}, '邮储银行')


def parse_jsbank(path, rows):
    """江苏银行：表头行2（序号/付款账号/交易日期/…/借贷标记/交易金额/账户余额）借贷标记借=支贷=收"""
    return _grow(path, rows, {'skip': 3, 'date': 2, 'dir': 5, 'inc': 6, 'exp': 6, 'bal': 7, 'cp': 1}, '江苏银行')


def parse_tailong(path, rows):
    """泰隆银行（3200 实证 2026-08-14）：表头行1（序号/交易时间/对方账号/对方户名/对方开户行/
    支出金额/收入金额/账户余额/摘要/附言），数据行2+；支出=支、收入=收。"""
    return _grow(path, rows, {'skip': 2, 'date': 1, 'exp': 5, 'inc': 6, 'bal': 7,
                              'cp': 3, 'sm': 8}, '泰隆银行')


def parse_pufa(path, rows):
    """浦发银行（3400 实证 2026-08-14）：表头行4（交易日期/交易时间/申请日期/凭证号/
    借方金额/贷方金额/余额/对方账号/对方户名/对方行名/交易流水号/摘要），数据行5+；
    借=支、贷=收；日期 YYYYMMDD。"""
    return _grow(path, rows, {'skip': 5, 'date': 0, 'dir': 0, 'exp': 4, 'inc': 5,
                              'bal': 6, 'cp': 8, 'sm': 11}, '浦发银行')


def parse_zheshang(path, rows):
    """浙商银行：表头行0（流水号/付款账号/付款人户名/收款账号/收款人户名/交易日期/交易金额/借贷标志(1收2付)/备注）。
    ⚡ 2026-08-13 修复：方向必须按『借贷标志』列（1=收/2=付）——原按『付款人户名』判断方向
    （_grow 的 zn 逻辑：付款人含『国望』→ 支），但浙商文件付款人户名恒为本方『国望』（本企业
    在工行/建行等多家行开立账户），他行国望→浙商国望的**收款**行付款人户名仍=国望 → 全部误判
    为支（3100 浙商实证：网银 2-04『支 1 亿』实为收 1 亿（工行国望转入），10 笔 2.78 亿方向全反，
    对应账侧『内转』借 2.78 亿全部错位；修复后配平）。"""
    hi, hdr = _probe_hdr(rows, ['流水号', '交易日期'])
    if hi is None:
        return []
    dc = _col(hdr, '交易日期')
    amtc = _col(hdr, '交易金额') or _col(hdr, '金额')
    dir_c = _col(hdr, '借贷标志')
    pc = _col(hdr, '付款人户名')
    rc = _col(hdr, '收款人户名')
    smc = _col(hdr, '备注')
    if dc is None or amtc is None or dir_c is None:
        return []
    out = []
    maxc = max(dc, amtc, dir_c, pc or 0, rc or 0, smc or 0)
    for r in rows[hi + 1:]:
        if not r or len(r) <= maxc:
            continue
        amt = _clean_amt(r[amtc])
        if amt <= 0:
            continue
        date = _norm_date(str(r[dc])[:12])
        if not date:
            continue
        flag = str(r[dir_c]).strip().replace('.0', '')
        direction = '收' if flag == '1' else ('支' if flag == '2' else ('收' if amt >= 0 else '支'))
        # 对方户名：本方付款(2)→收款人户名；本方收款(1)→付款人户名（⚠️ 收款人户名恒为国望时
        # 对方=付款人，如工行国望→浙商国望）
        cp = ''
        if direction == '支' and rc is not None:
            cp = str(r[rc]).strip()
        elif direction == '收' and pc is not None:
            cp = str(r[pc]).strip()
        out.append({
            'bank': '浙商', 'file': os.path.basename(path), 'account_no': '',
            'date': date, 'direction': direction, 'amount': amt,
            'counterparty': cp,
            'summary': str(r[smc]).strip()[:60] if smc is not None and len(r) > smc else '',
            'balance': None,
        })
    return out


def parse_suzhou(path, rows):
    """苏州银行：表头探测（账户交易明细表；交易日期/收入/支出/余额/对方户名）"""
    return parse_probed(path, rows, '苏州银行', ['交易日期', '收入'], '收入', '支出', '交易日期', '对方户名')


def parse_jiaxing(path, rows):
    """嘉兴银行（2026-08-13 DQ 账套新增）：表头行0（本方账号/对方账号/对方户名/交易时间/
    借/贷/借方发生额/贷方发生额/账户余额），数据行1+。
    银行角度：借=借方发生额=存款减少=支出；贷=贷方发生额=存款增加=收入。"""
    out = []
    acct = str(rows[0][0]).strip() if rows else ''
    for r in rows[1:]:
        if not r or len(r) < 8 or not str(r[3]).strip():
            continue
        d = _norm_date(r[3])
        if not d:
            continue
        inc = _clean_amt(r[6]) if len(r) > 6 else 0.0
        exp = _clean_amt(r[5]) if len(r) > 5 else 0.0
        if inc == 0 and exp == 0:
            continue
        out.append({
            'bank': '嘉兴银行', 'file': os.path.basename(path), 'account_no': acct,
            'date': d, 'direction': '收' if inc else '支',
            'amount': inc if inc else exp,
            'counterparty': str(r[2]).strip() if len(r) > 2 else '',
            'summary': '', 'balance': _clean_amt(r[7]) if len(r) > 7 else None,
        })
    return out


def parse_minsheng(path, rows):
    """民生银行：表头探测（交易/记账日期 + 借贷发生额），列名可能为收入/支出或借方/贷方"""
    hi, hdr = _probe_hdr(rows, ['交易'])
    if hi is None:
        hi, hdr = _probe_hdr(rows, ['日期', '金额'])
    if hi is None:
        return []
    dc = _col(hdr, '交易日期') or _col(hdr, '记账日期') or _col(hdr, '日期') or _col(hdr, '交易时间')
    ic = _col(hdr, '收入') or _col(hdr, '贷方')
    ec = _col(hdr, '支出') or _col(hdr, '借方')
    if dc is None or (ic is None and ec is None):
        return []
    return _grow(path, rows, {'skip': hi + 1, 'date': dc, 'inc': ic, 'exp': ec,
                              'cp': _col(hdr, '户名'), 'sm': _col(hdr, '摘要')}, '民生银行')


def _dispatch_icbc(path, rows):
    """工行：csv（GBK 制表符）走 parse_icbc_csv；xls/xlsx 若为『[HISTORYDETAIL] 单列表头』
    CSV 包装（3200 实证：凭证号/本方账号/对方账号/交易时间/借·贷/借方发生额/贷方发生额/…，
    与 parse_icbc_csv ①分支结构一致）也走 parse_icbc_csv——原 _dispatch 对 xlsx 一律走
    parse_icbc（期望借贷标志 col3/转出 col8/转入 col9）→ 3200 工行列错位解析出 6.79e21
    天文数字。非 HISTORYDETAIL 包装（3100 xls 标准表）走 parse_icbc。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        return parse_icbc_csv(path, rows)
    if rows and rows and str(rows[0][0] if rows[0] else '').strip() == '[HISTORYDETAIL]':
        return parse_icbc_csv(path, rows)
    return parse_icbc(path, rows)


def parse_boc_csv(path, rows):
    """中行 csv（UTF-8 BOM，逗号分隔）：行1 查询账号、行7 表头
    （交易类型/业务类型/付款人开户行号/…/付款人账号/付款人名称/…/收款人账号/收款人名称/
    交易日期/交易时间/交易货币/交易金额[带符号]/…/摘要/用途），行8+ 数据。
    金额：正=来账(收)、负=往账(支)。对方=收款人名称（本方付款时）否则付款人名称。"""
    qa = ''
    for r in rows[:3]:
        if any('查询账号' in str(x) for x in r):
            for x in r:
                s = str(x).strip().replace('\t', '')
                if s.isdigit() and len(s) >= 10:
                    qa = s
                    break
    hi = None
    for i, r in enumerate(rows):
        if any('交易类型' in str(x) and '交易金额' in str(x) for x in r):
            hi = i
            break
    if hi is None:
        for i, r in enumerate(rows):
            if any('交易类型' in str(x) for x in r):
                hi = i
                break
    if hi is None:
        return []
    out = []
    for r in rows[hi + 1:]:
        if len(r) < 14:
            continue
        try:
            amt = float(str(r[13]).replace(',', '').strip())
        except ValueError:
            continue
        if amt == 0:
            continue
        payer = str(r[5] or '').strip()
        payee = str(r[9] or '').strip()
        payer_acct = str(r[4] or '').strip().replace('\t', '')
        cp = payee if ((qa and payer_acct == qa) or not payee) else payer
        out.append({
            'bank': '中行', 'file': os.path.basename(path), 'account_no': qa,
            'date': _norm_date(str(r[10])[:14]), 'direction': '收' if amt > 0 else '支',
            'amount': abs(amt), 'counterparty': cp,
            'summary': (str(r[23] if len(r) > 23 else '') + ' ' +
                        str(r[24] if len(r) > 24 else '')).strip()[:40],
            'balance': None,
        })
    return [x for x in out if x.get('date')]


def _dispatch_boc(path, rows):
    """中行：csv（UTF-8 BOM）走 parse_boc_csv，xls/xlsx 走 parse_boc。"""
    ext = os.path.splitext(path)[1].lower()
    return parse_boc_csv(path, rows) if ext == '.csv' else parse_boc(path, rows)


def parse_ckb(path, rows):
    """进出口银行：表头（记账日期/借方发生额(支取)/贷方发生额(收入)/余额/币种/对方户名/对方账号/对方开户机构/业务发起时间/摘要）
    ⚡ 2026-08-14 修复：原写死 skip=1,date=2,exp=3,inc=4,cp=7 列映射全错——把贷方金额
    6550000.0 当日期（→6550-00-00）、余额 6,571,121.66 当支出、币种当收入、开户机构当
    对方户名 → 进出口差异 655 万假差异。改表头探测：记账日期0/借方1/贷方2/余额3/对方户名5。"""
    return parse_probed(path, rows, '进出口行', ['记账日期', '借方发生额'],
                        '贷方发生额', '借方发生额', '记账日期', cp_kw='对方户名')


# ---------------------------------------------------------------------------
# 注册表：按文件名关键词识别银行（关键词从长到短匹配，防 '中行' 误匹配 '国开行' 等）
# ---------------------------------------------------------------------------
BANK_RULES = [
    ('国开行', parse_cdb),
    ('国开', parse_cdb),
    ('农商银行', parse_nsbank),
    ('农商行', parse_nsbank),
    ('进出口', parse_ckb),
    ('上海银行', parse_shanghai),
    ('光大', parse_gd_eb),
    ('兴业', parse_cib),
    ('北京银行', parse_bob),
    ('华夏', parse_huaxia),
    ('南京银行', parse_njbank),
    ('宁波', parse_nbbank),
    ('平安', parse_pingan),
    ('招行', parse_cmb),
    ('邮储', parse_psbc),
    ('江苏银行', parse_jsbank),
    ('浙商', parse_zheshang),
    ('苏州银行', parse_suzhou),
    ('泰隆银行', parse_tailong),
    ('泰隆', parse_tailong),
    # ⚡ 2026-08-14 短关键词（3200/3300/3400 文件名无『银行』后缀，如『3300江苏1-6月人民币』）；
    #    放长关键词之后，_identify_bank 按最长匹配优先，无冲突
    ('江苏', parse_jsbank),
    ('苏州', parse_suzhou),
    ('上海', parse_shanghai),
    ('北京', parse_bob),
    ('浦发', parse_pufa),
    ('民生', parse_minsheng),
    ('嘉兴银行', parse_jiaxing),
    ('中信', parse_zhongxin),
    ('建行', parse_ccb),
    ('交通银行', parse_bocom),
    ('交行', parse_bocom),
    ('农业银行', parse_abc),
    ('农行', parse_abc),
    ('工行', _dispatch_icbc),
    ('工银', parse_icbc_asia),  # ⚡ 2026-08-14 3800『工银』=工银亚洲香港繁体（多账户多币种）
    ('中国银行', _dispatch_boc),
    ('中行', _dispatch_boc),
]


def _identify_bank(fname):
    """按文件名关键词识别银行；多关键词取最长匹配（防子串误判）。"""
    best, best_len = None, 0
    for kw, _fn in BANK_RULES:
        if kw in fname and len(kw) > best_len:
            best, best_len = kw, len(kw)
    return best


def parse_bank_file(path):
    """解析单个网银文件 → (bank, [统一行])；无法识别银行返回 (None, [])。"""
    fname = os.path.basename(path)
    bank = _identify_bank(fname)
    if bank is None:
        return None, []
    for kw, fn in BANK_RULES:
        if kw == bank:
            ext = os.path.splitext(fname)[1].lower()
            if ext == '.csv':
                rows = _rows_of_csv(path)
            else:
                rows = _rows_of_xls(path) if ext == '.xls' else _rows_of_xlsx(path)
            try:
                out = fn(path, rows)
                # ⚡ 2026-08-13 通用过滤：剔除合计/页脚行（'5 收 91,007,497.02' 等
                #   首列非日期被误解析成数据 → 网银凭空多出一笔大额 → 永不匹配）
                out = [r for r in out if r.get('date')]
                # ⚡ 2026-08-14 返回 parser 内部银行名（out[0]['bank']）而非文件名关键词——
                #   短关键词（'江苏'/'苏州'/'上海'/'泰隆'）与长关键词（'江苏银行'…）若返回关键词，
                #   同一银行拆成两组（3300 实证：'上海' 0 笔 + '上海银行' 200 元），与账侧对不上
                if out:
                    return out[0].get('bank') or bank, out
                return bank, []
            except Exception as e:
                print(f'  ⚠️ 解析失败 [{fname}]：{e}')
                return bank, []
    return bank, []


def parse_bank_folder(folder):
    """解析整个网银文件夹 → {'bank': [统一行]}，并列 log 无法识别的文件。"""
    result = {}
    unknown = []
    total = 0
    for f in sorted(os.listdir(folder)):
        if not f.lower().endswith(('.xls', '.xlsx', '.csv')):
            continue
        bank, rows = parse_bank_file(os.path.join(folder, f))
        if bank is None:
            unknown.append(f)
            continue
        result.setdefault(bank, []).extend(rows)
        total += len(rows)
        print(f'  ✓ {bank}: {f} → {len(rows)} 笔')
    return result, unknown, total


if __name__ == '__main__':
    import sys
    folder = sys.argv[1] if len(sys.argv) > 1 else os.path.join(P.DATA_ROOT, '300', '6100网银明细26年1-7月')
    banks, unknown, total = parse_bank_folder(folder)
    print(f'\n共解析 {total} 笔，银行 {len(banks)} 家')
    if unknown:
        print('未识别文件:', unknown)
    for b, rows in banks.items():
        inc = sum(r['amount'] for r in rows if r['direction'] == '收')
        exp = sum(r['amount'] for r in rows if r['direction'] == '支')
        print(f'  {b}: 收 {inc:,.2f} / 支 {exp:,.2f}（{len(rows)} 笔）')
