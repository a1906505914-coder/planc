# -*- coding: utf-8 -*-
"""函证回函批量识别（2026-08-16，铁律131f-ocr）：
扫描目录内 函证回函 照片/扫描PDF/Word → OCR → confirm_parser 结构化 → 汇总 Excel+JSON。

用法：
  python confirm_batch.py <目录> [--out 汇总.xlsx] [--json 结果.json]
"""
import sys, os, glob, json
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def scan_and_parse(src_dir, out_xlsx=None, out_json=None):
    import contract_reader as CR
    import confirm_parser as CP
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    exts = ('.pdf', '.docx', '.doc', '.xlsx', '.txt', '.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff')
    files = []
    for root, _, fs in os.walk(src_dir):
        for f in fs:
            if f.lower().endswith(exts) and not f.startswith('~$'):
                files.append(os.path.join(root, f))
    files.sort()
    if not files:
        print('未找到文件')
        return None

    results = []
    for i, f in enumerate(files):
        name = os.path.basename(f)
        try:
            text, fmt = CR.extract_text(f)
            if not text.strip():
                results.append({'file': name, 'ocr_chars': 0, 'status': 'OCR空'})
                print(f'[{i+1}/{len(files)}] {name[:40]} OCR空')
                continue
            r = CP.parse_confirm(text)
            r.update({'file': name, 'ocr_chars': len(text), 'status': 'OK'})
            results.append(r)
            flag = '⚠️' if r.get('conclusion') == '不相符' else '✅'
            print(f'[{i+1}/{len(files)}] {name[:40]} {flag} {r.get("conclusion","?")} {r.get("ref_no","")} {r.get("auditee","")[:14]}')
        except Exception as ex:
            results.append({'file': name, 'status': f'FAIL {ex}'})
            print(f'[{i+1}/{len(files)}] {name[:40]} FAIL {ex}')

    if out_json:
        json.dump(results, open(out_json, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(f'JSON: {out_json}')

    if out_xlsx:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '函证回函汇总'
        hdr = ['文件', '类型', '函证编号', '被审计单位', '结论', '差异说明',
               '盖章单位', '回函日期', '函证金额项目', 'OCR字数', '状态']
        ws.append(hdr)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='D9E2F3')
            c.alignment = Alignment(horizontal='center')
        for r in results:
            amts = '；'.join(f"{a['item']}={a['amount']:,.0f}" for a in r.get('amounts', []))
            ws.append([r.get('file', ''), r.get('type', ''), r.get('ref_no', ''), r.get('auditee', ''),
                       r.get('conclusion', ''), r.get('discrepancy', '')[:50], r.get('stamp_unit', ''),
                       r.get('reply_date', ''), amts, r.get('ocr_chars', ''), r.get('status', '')])
        for i, w in enumerate([30, 14, 18, 30, 8, 40, 24, 12, 30, 10, 10], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
        wb.save(out_xlsx)
        print(f'Excel: {out_xlsx}')

    n_ok = sum(1 for r in results if r.get('status') == 'OK')
    n_mm = sum(1 for r in results if r.get('conclusion') == '不相符')
    print(f'完成: {len(results)} 份, 识别 {n_ok}, 不相符 {n_mm}')
    return results


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else '.'
    out_x = None
    out_j = None
    if '--out' in sys.argv:
        out_x = sys.argv[sys.argv.index('--out') + 1]
    if '--json' in sys.argv:
        out_j = sys.argv[sys.argv.index('--json') + 1]
    if not out_x:
        out_x = os.path.join(src, '函证回函_识别汇总.xlsx')
    if not out_j:
        out_j = os.path.join(src, '函证回函_识别结果.json')
    scan_and_parse(src, out_x, out_j)
