"""收集合并集团底稿各科目『附注汇总』sheet → 附注汇总_合并.xlsx
用法：python footnote_collect.py <底稿目录> <输出文件>
遍历 <目录>/*审计底稿_*.xlsx，提取每个文件的『附注汇总』sheet（首个含附注的 sheet），
按报表科目顺序（sort_report_names）排列成一本工作簿。"""
import sys
import os
import glob
import openpyxl
from audit_common import sort_report_names


def collect(src_dir, out_path):
    files = sorted(glob.glob(os.path.join(src_dir, '*审计底稿_*.xlsx')))
    # 名称映射：文件名→科目名
    sheet_map = {}   # 科目名 → (文件路径, 文件显示名, sheet 数据)
    order = []
    for fp in files:
        base = os.path.basename(fp)
        if base.startswith('_备份_'):
            continue  # 备份文件不入附注汇总
        name = base.replace('审计底稿', '').replace('.xlsx', '').split('_')[0].strip()
        if not name:
            continue
        try:
            wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        except Exception:
            continue
        sn = next((s for s in wb.sheetnames if '附注' in s), None)
        if sn is None:
            wb.close()
            continue
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        wb.close()
        if not rows:
            continue
        # 同名科目多文件（如 _2026_生成 / _AH合并 双命名）：保留 mtime 最新者，避免旧文件按文件名排序后覆盖新文件
        if name in sheet_map and os.path.getmtime(fp) <= os.path.getmtime(sheet_map[name][0]):
            continue
        sheet_map[name] = (fp, base, rows)
        order.append(name)
    if not sheet_map:
        print('未找到任何附注汇总 sheet')
        return 0
    # 报表顺序排列
    ordered = sort_report_names(list(sheet_map))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ordered:
        _, base, rows = sheet_map[name]
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        print(f'  ✓ {name} ({base})')
    wb.save(out_path)
    print(f'完成: {len(ordered)} 个科目 → {out_path}')
    return len(ordered)


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else r'D:/底稿测试/AH/底稿/2026/合并集团'
    out = sys.argv[2] if len(sys.argv) > 2 else r'D:/底稿测试/AH/底稿/2026/附注汇总_合并.xlsx'
    collect(src, out)
