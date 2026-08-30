# -*- coding: utf-8 -*-
"""租赁合同台账读取器。
《租赁合同台账_<year>.xlsx》是租赁复核的唯一合同数据源（人可读/程序可读），
含本次案例固化的全部因素：合同起止日/完整租赁期/付款周期/付款时点/完整期付款额/
免租形态/含税口径/物业费剥离等。程序读台账生成 contract 配置，不再依赖硬编码。

用法：
    from lease_ledger import contracts_for
    rows = contracts_for('GFY', '01FY本级')   # → [contract_dict, ...] 或 None（无台账/无该主体）
"""
import os
import sys

import paths as P

DATA_ROOT = P.DATA_ROOT  # 引用 paths 配置中心（AUDIT_DATA_ROOT 环境变量可覆盖）

FREQ_MAP = {'年付': 'year', '半年付': 'half', '季付': 'quart', '月付': 'month'}
FREQ_NUM = {'year': 12.0, 'half': 6.0, 'quart': 3.0, 'month': 1.0}


def _ledger_path(acct, year):
    """台账路径：优先 <数据根>/<acct>/数据/<year>/使用权资产租赁合同/租赁合同台账_<year>.xlsx。"""
    return os.path.join(DATA_ROOT, acct, '数据', str(year), '使用权资产租赁合同',
                        f'租赁合同台账_{year}.xlsx')


def _parse_pmts(s):
    """'1,234.56;7,890.00' → [1234.56, 7890.0]"""
    out = []
    for part in str(s or '').replace(',', '').replace('，', '').split(';'):
        part = part.strip()
        if not part:
            continue
        try:
            out.append(float(part))
        except ValueError:
            continue
    return out


def contracts_for(acct, sub, year=None):
    """返回台账中该主体的合同配置列表（与 lease_manifest contracts 结构一致）；
    无台账文件或该主体无合同行 → 返回 None（调用方回退 manifest）。"""
    try:
        import openpyxl
    except ImportError:
        return None
    if year is None:
        year = 2026
    fp = _ledger_path(acct, year)
    if not os.path.exists(fp):
        return None
    try:
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    except Exception:
        return None
    if '合同台账' not in wb.sheetnames:
        return None
    ws = wb['合同台账']
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None
    hdr = [str(x or '').strip() for x in rows[0]]
    try:
        i_sub = hdr.index('主体')
        i_name = hdr.index('合同名称/资产')
        i_sd = hdr.index('租赁开始日')
        i_ed = hdr.index('租赁结束日')
        i_pm = hdr.index('租赁期(月)')
        i_freq = hdr.index('付款周期')
        i_adv = hdr.index('付款时点')
        i_pay = hdr.index('每期付款额(分号分隔)')
        i_note = hdr.index('备注')
    except ValueError:
        return None
    out = []
    for r in rows[1:]:
        if r[i_sub] is None or str(r[i_sub]).strip() != sub:
            continue
        name = str(r[i_name] or '').strip()
        if not name:
            continue
        pmts = _parse_pmts(r[i_pay])
        if not pmts:
            continue
        freq = FREQ_MAP.get(str(r[i_freq] or '').strip(), 'year')
        adv = '期初' in str(r[i_adv] or '')
        try:
            pm = int(r[i_pm] or 0)
        except (TypeError, ValueError):
            pm = round(len(pmts) * FREQ_NUM.get(freq, 12.0))
        c = {
            'name': name,
            'start_date': str(r[i_sd] or '').strip()[:10],
            'end_date': str(r[i_ed] or '').strip()[:10],
            'period_months': pm,
            'payment_freq': freq,
            'advance': adv,
            'payments': pmts,
        }
        if r[i_note] not in (None, ''):
            c['remark'] = str(r[i_note])
        out.append(c)
    return out or None


if __name__ == '__main__':
    for sub in ('01FY本级', '05FY皮革'):
        rows = contracts_for('GFY', sub)
        if rows is None:
            print(sub, ': 无台账')
        else:
            print(f'=== {sub}: {len(rows)} 合同 ===')
            for c in rows:
                print(' ', c['name'], c['start_date'], c['end_date'],
                      'pm=', c['period_months'], c['payment_freq'],
                      'adv=', c['advance'], '期数', len(c['payments']))
