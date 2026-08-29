# -*- coding: utf-8 -*-
"""mask_dict.py —— 通用主体脱敏字典（S1 级，2026-08-19）。

跨专项统一：公司名/客商名/户名/合同方 → A 编号代码。
映射持久化到 mask_dict.json，同主体在【所有专项】中永远是同一代码（可交叉引用、可复现）。
真实名↔代码对应表仅保存在本地（审计人员可还原），对外只发脱敏版。

用法：
    from mask_dict import build, lookup
    build()                     # 增量重建（从 AH 公司代码等来源收集，追加新名字）
    build(extra_names={'xxx'})  # 补充额外主体名
    lookup('杭氧集团股份有限公司')  # → 'A001'；未收录返回 None（引擎保留原样）

扩展：新增数据源 = 在 _collect_sources() 里加一个收集器（返回真实名集合），
      如科目余额表主体、开户清单存款人、合同出租/承租方等。
"""
import json
import os
import re

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DICT_FP = os.path.join(APP_DIR, 'mask_dict.json')

# 主体名特征（底稿 S1 列值过滤）
_COMPANY_PAT = re.compile(r'(有限公司|有限责任公司|股份有限公司|集团有限公司|集团公司|股份公司|'
                          r'公司|集团|工厂|厂|合伙企业|事务所|银行|合作社|分公司|总厂)$')

# 集团内【主体】列名关键词（我方/核算主体/公司名等；不含客户/供应商/客商——外部客商走【客商】兜底）
_S1_COL_KEYS = ('核算主体', '我方主体', '我方名称', '主体名称', '公司名称', '单位名称',
                '账户名称', '存款人', '承租方', '出租方', '借款方', '贷款方',
                '本公司', '集团内', '关联方名称')
# 明确【外部客商】列名（不收集进字典，引擎 S1 兜底【客商】）
_EXCLUDE_COL_KEYS = ('客户', '供应商', '客商', '往来单位', '对方单位', '对方名称',
                     '对方主体', '对手', '付款方', '收款方', '交易对手')


def _collect_sources(full=False):
    """收集全项目真实主体名（增量来源，可扩展）。
    full=True 额外遍历各账套底稿 xlsx 的【集团内主体】列（慢，首次建字典用）。
    外部客商不收集（数量大且无交叉引用价值），由引擎 S1 兜底【客商】。"""
    names = set()
    # ① AH 公司代码.xlsx（137 家，关联往来/序时账主体）
    try:
        from ah_intra_seq_extract import load_company_map
        names |= set(load_company_map().values())
    except Exception:
        pass
    # ② AZ 回函映射（兆龙系全名，confirm_mask.json auditee_codes）
    try:
        cm_fp = os.path.join(APP_DIR, 'confirm_mask.json')
        if os.path.exists(cm_fp):
            with open(cm_fp, encoding='utf-8') as f:
                cm = json.load(f)
            names |= set(cm.get('auditee_codes', {}).keys())
    except Exception:
        pass
    # ③ 借款合同主体（contract_mask.json borrower_dir 目录名，如"3100国望"→取"国望"）
    try:
        ct_fp = os.path.join(APP_DIR, 'contract_mask.json')
        if os.path.exists(ct_fp):
            with open(ct_fp, encoding='utf-8') as f:
                ct = json.load(f)
            for d in ct.get('borrower_dir', {}):
                m = re.search(r'\d+(.+)', d)
                if m:
                    names.add(m.group(1).strip())
    except Exception:
        pass
    # ④ 全账套底稿集团内主体列（full 才跑，慢）
    if full:
        try:
            names |= _collect_from_books()
        except Exception:
            pass
    return names


def _collect_from_books():
    """遍历各账套【底稿/中间产物/prepared】xlsx，从【集团内主体】列收集主体名。
    严格过滤：整格=公司名（无标点/数字、以公司后缀结尾），排除外部客商列。"""
    import openpyxl
    names = set()
    root_dir = os.environ.get('AUDIT_DATA_ROOT', r'D:\底稿测试')
    for root, dirs, files in os.walk(root_dir):
        if not any(k in root for k in ('底稿', 'prepared', '中间产物')):
            continue
        for f in files:
            if not f.endswith('.xlsx') or f.startswith('~$'):
                continue
            fp = os.path.join(root, f)
            try:
                wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
                try:
                    for sn in wb.sheetnames[:10]:
                        ws = wb[sn]
                        hdr = None
                        for r in ws.iter_rows(values_only=True, max_row=10):
                            cells = [str(c or '').strip() for c in r]
                            if any(c and any(k in c for k in _S1_COL_KEYS) for c in cells):
                                hdr = cells
                                break
                        if not hdr:
                            continue
                        s1_idx = [i for i, h in enumerate(hdr)
                                  if any(k in h for k in _S1_COL_KEYS)
                                  and not any(k in h for k in _EXCLUDE_COL_KEYS)]
                        for r in ws.iter_rows(values_only=True, max_row=500):
                            for i in s1_idx:
                                if i < len(r) and r[i]:
                                    v = str(r[i]).strip()
                                    # 整格必须是纯公司名：无标点/数字/换行，以公司后缀结尾
                                    if (4 <= len(v) <= 40
                                            and not re.search(r'[，。；、：:""（）()\s\d]', v)
                                            and _COMPANY_PAT.search(v)):
                                        names.add(v)
                finally:
                    wb.close()
            except Exception:
                continue
    return names


def _load():
    if os.path.exists(DICT_FP):
        try:
            with open(DICT_FP, encoding='utf-8') as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def _save(data):
    with open(DICT_FP, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=1)


def build(extra_names=None, force=False, full=False):
    """重建/增量字典。force=True 清空重建（编号会变，慎用）；
    full=True 额外遍历各账套底稿 S1 列收集主体（慢，首次建字典用）。
    返回 (data, added)。已有映射保留，只追加新名字（编号稳定）。"""
    data = {} if force else _load()
    names = _collect_sources(full=full)
    if extra_names:
        names |= set(extra_names)
    existing = set(data)
    added = 0
    for nm in sorted(names - existing):
        if not nm or nm in ('', '合计', '小计'):
            continue
        if not re.search(r'[\u4e00-\u9fff]', nm):
            continue  # 过滤无中文字符的噪音条目（如公司代码表头 "SAP A.G."）
        data[nm] = 'A%03d' % (len(data) + 1)
        added += 1
    if added or force:
        _save(data)
    return data, added


def lookup(name):
    """真实名 → 代码；未收录返回 None。"""
    if not name:
        return None
    return _load().get(name)


def mask_name(name, fallback='【客商】'):
    """主体名 → 脱敏代码（字典命中 A 代码；未收录含公司特征 → 兜底词；否则原样）。
    供程序 console 打印前脱敏，防止主体名进对话/日志。"""
    if not name:
        return name
    code = lookup(str(name).strip())
    if code:
        return code
    if re.search(r'(有限公司|公司|集团|工厂|厂|合伙企业|事务所|银行|合作社|分公司)', str(name)):
        return fallback
    return name


def mask_names(names, maxn=20):
    """主体名集合/列表 → 脱敏代码列表（批量打印用）。"""
    out = [mask_name(n) for n in names]
    if maxn and len(out) > maxn:
        out = out[:maxn] + [f'…共{len(out)}家']
    return out


def names():
    """返回 {真实名: 代码} 全量映射（引擎 S1sub 用）。"""
    return _load()


if __name__ == '__main__':
    import sys
    full = '--full' in sys.argv
    force = '--force' in sys.argv
    data, added = build(full=full, force=force)
    print('字典主体数: %d | 本次新增: %d' % (len(data), added))
    for k in list(data)[:6]:
        print('  %s = %s' % (data[k], k))
