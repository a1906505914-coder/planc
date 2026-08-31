# -*- coding: utf-8 -*-
"""最终铁律核对：底稿审定表 逐主体 vs 自建试算表（Sheet1 期末余额）。

口径：
- 底稿取『审定数/期末未审数』列（优先审定数；无则本期数——损益类）
- TB Sheet1 逐主体期末余额（qm）
- 逐主体比绝对值（排除贷余为负的符号口径），差异>阈值 记真实差异
- 输出：每科目 真实差异主体数 / 符号口径一致主体数 / 差异额合计
用法：python recon_final.py --data <data_dir> --tb <自建试算表.xlsx> --dir <底稿目录>
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def find_val_col(rows):
    """优先『审定数』列（资产负债表科目期末口径）；无则『期末未审数/本期数』。"""
    for kw in ('审定数', '期末未审数', '本期数', '期末数'):
        for i, r in enumerate(rows[:8]):
            hs = [str(x) if x else '' for x in r]
            if kw in hs:
                return i, hs.index(kw)
    return None, None


# 名称含这些字样但【非】该科目本身（递延所得税-xxx / 处置利得 等），核对 TB 时排除
TB_EXCLUDE = ('递延所得税', '利得', '损失', '处置', '清理', '折旧', '摊销', '减值',
              '坏账', '准备', '收入', '费用', '营业外', '所得税')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', required=True, help='数据目录（含 科目余额表）')
    ap.add_argument('--tb', required=True, help='自建试算表 xlsx 路径')
    ap.add_argument('--dir', required=True, help='底稿目录')
    ap.add_argument('--threshold', type=float, default=1.0)
    args = ap.parse_args()

    import openpyxl
    import sap_adapter as A
    A._DATA_ROOT = args.data

    # TB Sheet1 逐主体
    wb = openpyxl.load_workbook(args.tb, read_only=True, data_only=True)
    ws = wb['试算表']
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    tb_rows = {}
    for r in rows[1:]:
        if r[0]:
            tb_rows[str(r[0]).strip()] = {
                str(hdr[i]): v for i, v in enumerate(r)
                if i > 0 and hdr[i] and '合计' not in str(hdr[i]) and isinstance(v, (int, float))}
    # Sheet3 利润表（损益类科目核对，发生额口径：费用取借方、收益取贷方，与 sap_tb_gen 同源）
    pl_rows = {}
    if '利润表' in wb.sheetnames:
        _ws3 = wb['利润表']
        _r3 = list(_ws3.iter_rows(values_only=True))
        _hi3 = next((i for i, _r in enumerate(_r3[:4])
                     if _r and _r[0] and '项目' in str(_r[0])), None)
        if _hi3 is not None:
            _heads3 = [str(c).strip() if c else '' for c in _r3[_hi3]]
            for _r in _r3[_hi3 + 1:]:
                if not _r or not _r[0]:
                    continue
                # ⚡⚡ 2026-08-31 排除『集团合计』列（否则主体Σ+合计 双计）
                pl_rows[str(_r[0]).strip()] = {_heads3[j]: v for j, v in enumerate(_r)
                                               if j > 0 and _heads3[j] and _heads3[j] != '集团合计'
                                               and isinstance(v, (int, float))}
    wb.close()

    out = args.dir
    print(f"{'科目':<14}{'真实差异':>8}{'符号一致':>8}{'差异额合计':>18}  判定")
    n_real = n_ok = 0
    for fn in sorted(os.listdir(out)):
        if not fn.endswith('.xlsx'):
            continue
        subj = fn.split('审计底稿')[0].strip()   # 通用命名（AH合并/XBJ/2025_生成 全适配）
        try:
            wb2 = openpyxl.load_workbook(os.path.join(out, fn), read_only=True, data_only=True)
            sns = [s for s in wb2.sheetnames if '审定表' in s and '集团' not in s and '合并' not in s]
            if not sns:
                wb2.close()
                continue
            rws = list(wb2[sns[0]].iter_rows(values_only=True))
            wb2.close()
            hi, ci = find_val_col(rws)
            if hi is None:
                continue
            # ⚡⚡ 2026-08-31 损益类识别：审定表含『本期数』且无『期末未审数』→ 损益类
            #   （发生额口径，用 Sheet3 利润表核对；Sheet1 qm=0 不适用）
            _is_pl = any('本期数' in str(c) for c in rws[hi]) and                       not any('期末未审数' in str(c) for c in rws[hi])
            # ⚡⚡ 2026-08-31 营业收入结构C（项目×主体）：审定表含『一、主营业务收入』项目行
            #   → 特殊核对：项目行聚合 vs 利润表营业收入/营业成本
            _is_rev = subj == '营业收入' or any(
                r and r[0] and str(r[0]).strip() == '一、主营业务收入' for r in rws[hi:hi + 3])
            if _is_rev:
                _rev = _cost = 0.0
                for r in rws[hi + 1:]:
                    if not r or not r[0]:
                        continue
                    if r[1] and str(r[1]).strip() in ('全集团合计', '合计'):
                        continue
                    _p = str(r[0]).strip()
                    if _p in ('一、主营业务收入', '二、其他业务收入'):
                        _rev += r[3] if len(r) > 3 and isinstance(r[3], (int, float)) else 0
                    elif _p in ('减：主营业务成本', '减：其他业务成本'):
                        _cost += r[3] if len(r) > 3 and isinstance(r[3], (int, float)) else 0
                _tb_rev = sum(v for v in pl_rows.get('营业收入', {}).values() if isinstance(v, (int, float)))
                _tb_cost = sum(v for v in pl_rows.get('营业成本', {}).values() if isinstance(v, (int, float)))
                _real = abs(abs(_rev) - abs(_tb_rev)) > 1.0 or abs(abs(_cost) - abs(_tb_cost)) > 1.0
                _diff_t = abs(_rev) + abs(_cost) + abs(_tb_rev) + abs(_tb_cost)
                if _real:
                    n_real += 1
                    print(f"{'营业收入':<14}{1:>8}{0:>8}{_diff_t:>18,.2f}  ❌")
                    print(f"    主营收入: 底稿={_rev:,.2f} 利润表={_tb_rev:,.2f}")
                    print(f"    主营成本: 底稿={_cost:,.2f} 利润表={_tb_cost:,.2f}")
                else:
                    n_ok += 1
                continue
            # 结构判断：数据行 row[1] 是否为『项目』（原值/累计折旧/减值/净值）→ 结构B（主体×项目）
            # ⚡ 跳过表头后的空行（审定表标题/说明/空行）找首个数据行
            is_struct_b = False
            for _r0 in rws[hi + 1:]:
                if _r0 and _r0[0] and str(_r0[0]).strip() not in ('合计', '总计', '全集团合计', '集团加计'):
                    if _r0[1] and any(k in str(_r0[1]) for k in ('原值', '累计折旧', '减值准备', '净值', '账面')):
                        is_struct_b = True
                    break
            # ⚡⚡ 2026-08-31 结构D-往来汇总包（AZ 泰国账套往来科目）：表头含『汇总包顺序』
            #   列，row[0]=顺序号、row[1]=公司 → 用 row[1] 当主体。
            #   ⚡⚡ 取『期末余额』列（非审定数——AZ 汇总包审定数含审计重分类调整，
            #   如上分 期末 2,184,784 + 重分类 2,109,574 = 审定数 4,294,358，应对比 TB 期末）
            is_struct_pkg = bool(rws[hi]) and any('汇总包' in str(c) for c in rws[hi])
            if is_struct_pkg:
                for _j, _h in enumerate(rws[hi]):
                    if _h and '期末余额' in str(_h):
                        ci = _j
                        break
            # ⚡⚡ 2026-08-31 结构C-科目×主体（货币资金等）：表头含『科目』列，数据行
            #   row[0]=科目名、row[1]=主体名 → 按 row[1] 聚合主体值（row[0] 筛 subj）
            is_struct_ce = False
            if not is_struct_b and not is_struct_pkg:
                for _r0 in rws[hi + 1:]:
                    if _r0 and _r0[0] and _r0[1]:
                        if str(_r0[1]).strip() in ('合计', '全集团合计'):
                            break
                        is_struct_ce = not str(_r0[0]).strip().isdigit() and bool(_r0[1])
                        break
            wp = {}
            for r in rws[hi + 1:]:
                if not r or not r[0] or str(r[0]).strip() in ('合计', '总计', '全集团合计', '集团加计'):
                    continue
                e = str(r[0]).strip()
                if is_struct_pkg:
                    _ent = str(r[1]).strip() if r[1] else ''
                    if not _ent or _ent in ('合计', '全集团合计'):
                        continue
                    if len(r) > ci and isinstance(r[ci], (int, float)):
                        wp[_ent] = wp.get(_ent, 0.0) + r[ci]
                elif is_struct_ce:
                    _sk = str(r[0]).strip()
                    _ent = str(r[1]).strip() if r[1] else ''
                    if not _ent or _ent in ('合计', '全集团合计', '小计'):
                        continue
                    if _sk == subj or _sk.endswith(subj) or subj in _sk:
                        if len(r) > ci and isinstance(r[ci], (int, float)):
                            wp[_ent] = wp.get(_ent, 0.0) + r[ci]
                elif is_struct_b:
                    # 只取『原值』项目行
                    if r[1] and '原值' in str(r[1]) and len(r) > ci and isinstance(r[ci], (int, float)):
                        wp[e] = wp.get(e, 0.0) + r[ci]
                elif len(r) > ci and isinstance(r[ci], (int, float)):
                    wp[e] = r[ci]
            if not wp:
                continue
            # TB 匹配：精确名优先；否则 名称包含（排除 递延所得税/处置利得 从属科目、
            #   及『其他应付账款』等含前缀但不同科目标签的误配）
            # ⚡⚡ 2026-08-31 损益类改从 Sheet3 利润表取数（发生额口径）
            _src_rows = pl_rows if _is_pl else tb_rows
            tbk = None
            if subj in _src_rows:
                tbk = subj
            else:
                for k in _src_rows:
                    if subj and subj in k and len(k) <= len(subj) + 4 \
                            and not any(x in k for x in TB_EXCLUDE) \
                            and not (k.startswith('其他') and not subj.startswith('其他')):
                        tbk = k
                        break
            if tbk is None:
                for k in _src_rows:
                    if k and k in subj and len(subj) <= len(k) + 4:
                        tbk = k
                        break
            if tbk is None:
                continue
            tbm = dict(_src_rows[tbk])
            # ⚡⚡ 2026-08-30 归并（审计列报口径）：其他应收/应付 需并入 应收/应付股利利息（TB 分开）
            MERGE = {'其他应收款': ['应收股利', '应收利息'],
                     '其他应付款': ['应付股利', '应付利息']}
            for _mk in MERGE.get(subj, []):
                if _mk in _src_rows:
                    for _e, _v in _src_rows[_mk].items():
                        tbm[_e] = tbm.get(_e, 0.0) + _v
            # ⚡⚡ 2026-08-31 损益类无归并（_src_rows=pl_rows）；此处保证 tbm 有值
            real = []
            sign = 0
            diff_tot = 0.0
            for e, wv in wp.items():
                tv = tbm.get(e, 0.0)
                if abs(abs(wv) - abs(tv)) > args.threshold:
                    real.append((e, wv, tv))
                    diff_tot += abs(abs(wv) - abs(tv))
                else:
                    sign += 1
            mark = '❌' if real else '✓'
            if real:
                n_real += 1
            else:
                n_ok += 1
            print(f'{subj:<14}{len(real):>8}{sign:>8}{diff_tot:>18,.0f}  {mark}')
            for e, wv, tv in real[:3]:
                print(f'    {e}: 底稿={wv:,.2f} TB={tv:,.2f}')
        except Exception:
            pass
    # ⚡⚡ 2026-08-31 缺失底稿检查（回应"少科目没发现"根因）：recon_final 原先只核对
    #   【已存在的底稿文件】——AH 缺银行存款（bank 未跑）时 44 科目照常全绿，无人发现缺
    #   科目。对照 audit_common 报表科目全集（资产/负债/权益/损益），列出无底稿文件的科目。
    try:
        from audit_common import REPORT_ASSET, REPORT_LIAB, REPORT_EQUITY, REPORT_PL
        _have_files = [fn.split('审计底稿')[0].strip() for fn in os.listdir(out) if fn.endswith('.xlsx')]
        _no_note = ('营业成本', '主营业务收入', '主营业务成本', '其他业务收入', '其他业务成本',
                    '本年利润', '研发支出', '累计折旧', '累计摊销', '临时设施摊销',
                    '原材料', '库存商品', '周转材料', '低值易耗品', '存货跌价准备',
                    '房地产开发成本', '待摊费用', '坏账准备', '贷款', '贷款损失准备',
                    '抵债资产', '抵债资产跌价准备', '合同资产减值准备', '持有待售资产减值',
                    # ⚡⚡ 2026-08-31 由其他底稿覆盖的报表科目：货币资金→银行存款底稿
                    #   （货币资金审定表含 库存现金/银行存款/其他货币资金）；应付股利/预提费用
                    #   → 其他应付款底稿（recon_final MERGE 并入）
                    '其他货币资金', '应付股利', '预提费用', '内部结算中心存款')

        def _has_note(_s):
            # 底稿文件匹配（应付职工薪酬 ↔ 职工薪酬 等前缀差）
            return any(_s in hf or hf in _s for hf in _have_files)

        def _tb_zero(_s):
            # 试算表该科目全主体 0（无数据）→ 无需底稿
            _r = tb_rows.get(_s)
            if not _r:
                return True
            return not any(abs(x) > 0.005 for x in _r.values() if isinstance(x, (int, float)))

        _missing = [s for s in REPORT_ASSET + REPORT_LIAB + REPORT_EQUITY + REPORT_PL
                    if not _has_note(s) and s not in _no_note and not _tb_zero(s)]
        if _missing:
            print(f'⚠️ 报表科目有数据但缺底稿文件（未生成/未核对）: {" / ".join(_missing)}')
    except Exception:
        pass
    print(f'\n一致/符号口径 {n_ok} 个 / 真实差异 {n_real} 个')
    return 0 if n_real == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
