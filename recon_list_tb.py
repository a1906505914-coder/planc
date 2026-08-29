# -*- coding: utf-8 -*-
"""往来清单 vs TB 口径核对（按科目码，抽查主体）。
对指定主体：读 应付/{comp}.xlsx / 应收/{comp}.xlsx 清单，按科目码聚合期末余额；
与 read_km(comp) 的 TB 期末对比，展示符号与金额差异规律。
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl, collections
import sap_adapter as A

DATA = r'd:/底稿测试/AH/数据/2026'
A.set_root(DATA)

def read_list(comp, kind):
    """kind: 'ap' 应付 / 'ar' 应收。返回 {科目码前缀6: 期末合计, 行数}"""
    dirname = '应付' if kind == 'ap' else '应收'
    fp = os.path.join(DATA, dirname, f'{comp}.xlsx')
    if not os.path.exists(fp):
        return None
    wb = openpyxl.load_workbook(fp, data_only=True, read_only=False)
    ws = wb[wb.sheetnames[0]]
    agg = collections.defaultdict(float)
    n = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[3]:
            continue
        code = str(r[3]).strip()
        try:
            qm = float(r[9] or 0)
        except (TypeError, ValueError):
            qm = 0.0
        agg[code[:6]] += qm
        n += 1
    wb.close()
    return agg, n

def tb_codes(comp):
    try:
        km = A.read_km(comp)
    except Exception:
        return {}
    out = collections.defaultdict(float)
    for c, r in km.items():
        if abs(float(r.get('qm') or 0)) > 0.005:
            out[str(c)[:6]] += float(r.get('qm') or 0)
    return out

def show(comp, kinds=('ap', 'ar')):
    print(f'===== 主体 {comp} =====')
    for kind in kinds:
        res = read_list(comp, kind)
        if res is None:
            print(f'  {kind}: 无清单文件')
            continue
        agg, n = res
        tb = tb_codes(comp)
        print(f'  -- {kind} 清单 {n} 行，科目码{len(agg)}个 | TB 科目码{len(tb)}个 --')
        # 清单科目码与 TB 对齐（前缀6位）：只对比清单里出现的码（清单仅含往来科目）
        codes = sorted(set(agg))
        shown = 0
        for code in codes:
            lv = agg.get(code, 0.0)
            tv = tb.get(code, 0.0)
            if abs(lv) < 1 and abs(tv) < 1:
                continue
            same_sign = (lv >= 0) == (tv >= 0) if (abs(lv) > 1 and abs(tv) > 1) else '?'
            ratio = (abs(lv) / abs(tv)) if abs(tv) > 1 else ('清单有TB无' if abs(lv) > 1 else '')
            flag = 'OK' if (isinstance(ratio, float) and 0.95 < ratio < 1.05 and same_sign is True) else ''
            print(f'    {code}: 清单={lv:>16,.2f} | TB={tv:>16,.2f} | 比={ratio} | {flag}')
            shown += 1
            if shown >= 12:
                break

if __name__ == '__main__':
    comps = sys.argv[1:] or ['1010']
    for c in comps:
        show(c)
