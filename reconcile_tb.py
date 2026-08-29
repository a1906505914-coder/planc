# -*- coding: utf-8 -*-
"""审计底稿 ↔ 科目余额表(TB) 勾稽复核报告。
两条独立复核线：
  (A) 资产负债表类【审定表】：期末未审数(按方向 signed) 与 TB 期末余额(signed qm) 比对。
      审定表由 add_audit_summary 直接自 TB 填充，故应严格相等；此处作为独立复核，
      捕捉任何科目映射/聚合偏差。
  (B) 利润表科目：应用「贷方蓝字→借方红字」规则后的 GL 自然方发生额 与 TB 自然方发生额
      比对，应相等（即等于结转损益口径）。这是本次方法论修正的核心验证。
      注：2026 GL 仅含 1–5 月(YTD)，而 2026 TB 为全年，故 2026 利润表行单独标注「YTD不可比」，
      不作为"不一致"误报。
输出：审计底稿_TB勾稽复核报告.csv + 控制台汇总。
"""
import paths as P
import os
import re
import csv
import openpyxl
import audit_common as A
import pl_detail as PL

DESKTOP = P.DATA_ROOT
FOLDERS = ["c", "J", "x", "T", "S"]
TOL = 0.01
YTD_YEARS = {'2026'}


def parse_codes(cell):
    if cell is None:
        return []
    s = str(cell).strip()
    if not s:
        return []
    parts = re.split(r'[;/,，、\s]+', s)
    return [p for p in parts if re.match(r'^\d{3,8}$', p)]


def find_header(ws):
    for ri in range(1, 9):
        row = [ws.cell(ri, ci).value for ci in range(1, min(ws.max_column, 16) + 1)]
        if any('科目编码' in str(x) for x in row if x):
            return ri, row
    return None, None


def col_index(header, name):
    for i, h in enumerate(header):
        if h and name in str(h):
            return i + 1
    return None


def signed_qm(tb_full, codes, year):
    s = 0.0
    for (e, c, n, y), v in tb_full.items():
        if str(y) == str(year) and c in codes:
            s += v['qm']
    return s


# ---------------- (A) BS 审定表 ----------------
def reconcile_bs(fp, tb_full):
    wb = openpyxl.load_workbook(fp, data_only=True)
    out = []
    for ws in wb.worksheets:
        if '审定表' not in ws.title:
            continue
        hri, header = find_header(ws)
        if not header:
            out.append((ws.title, '—', '—', 0.0, 0.0, 'NO_HEADER'))
            continue
        code_col = col_index(header, '科目编码')
        year_col = col_index(header, '年度')
        dir_col = col_index(header, '期末方向')
        val_col = col_index(header, '期末未审数')
        if not (code_col and year_col and dir_col and val_col):
            out.append((ws.title, '—', '—', 0.0, 0.0, 'NO_COL'))
            continue
        for ri in range(hri + 1, ws.max_row + 1):
            codes = parse_codes(ws.cell(ri, code_col).value)
            if not codes:
                continue
            year = ws.cell(ri, year_col).value
            dirc = ws.cell(ri, dir_col).value
            val = ws.cell(ri, val_col).value
            if val is None:
                continue
            try:
                val = float(val)
            except Exception:
                continue
            sign = 1.0 if (dirc and '借' in str(dirc)) else (-1.0 if (dirc and '贷' in str(dirc)) else 1.0)
            rep = val * sign
            exp = signed_qm(tb_full, codes, year)
            diff = rep - exp
            ok = abs(diff) < TOL
            out.append((ws.title, ','.join(codes), str(year),
                        round(rep, 2), round(exp, 2), 'OK' if ok else 'DIFF %.2f' % diff))
    wb.close()
    return out


# ---------------- (B) 利润表 GL自然 == TB自然（按一级科目汇总比对） ----------------
def reconcile_pl(data_dir):
    ents = PL._discover_entities(data_dir)
    if not ents:
        return []
    years = sorted({y for b in ents.values() for y in b})
    gl_agg, _ = PL.read_pl_gl(data_dir, ents, years)
    tb = PL.read_tb_control(data_dir, ents, years)
    # 将二级 GL drill-down 汇总到一级(subj, e, y)，与 TB 一级控制数比对
    gl_nat = {}   # (subj, e, y) -> 自然方发生额
    for (subj, e, l2, y), (dr, cr) in gl_agg.items():
        kind = PL._pl_kind(subj)              # 'rev'(贷自然) / 'exp'(借自然)
        nat_gl = cr if kind == 'rev' else dr
        key = (subj, e, y)
        gl_nat[key] = gl_nat.get(key, 0.0) + nat_gl
    out = []
    for (subj, e, y), gv in sorted(gl_nat.items()):
        t = tb.get((e, subj, y))
        if not t:
            out.append((subj, e, y, '(一级汇总)', round(gv, 2), 0.0, 'NO_TB'))
            continue
        nat_tb = t[1] if PL._pl_kind(subj) == 'rev' else t[0]   # (jf, df)
        diff = gv - nat_tb
        ytd = (str(y) in YTD_YEARS)
        if ytd and abs(diff) >= TOL:
            concl = 'YTD不可比(2026 GL=1-5月)'
        else:
            concl = 'OK' if abs(diff) < TOL else 'DIFF %.2f' % diff
        out.append((subj, e, y, '(一级汇总)', round(gv, 2), round(nat_tb, 2), concl))
    return out


def main():
    rows_bs, rows_pl = [], []
    for f in FOLDERS:
        d = os.path.join(DESKTOP, f)
        if not os.path.isdir(d):
            continue
        ents = A.discover_entities(d)
        tb_full = A.read_tb_full(d, ents)
        for fn in sorted(os.listdir(d)):
            if not fn.endswith('_生成.xlsx') or '_bak' in fn:
                continue
            fp = os.path.join(d, fn)
            try:
                rows_bs += [(f, fn, *r) for r in reconcile_bs(fp, tb_full)]
            except Exception as ex:
                rows_bs.append((f, fn, 'ERR', '—', '—', 0.0, 0.0, str(ex)[:80]))
        try:
            rows_pl += [(f, *r) for r in reconcile_pl(d)]
        except Exception as ex:
            rows_pl.append((f, 'ERR', '—', '—', '—', 0.0, 0.0, str(ex)[:80]))

    outp = os.path.join(DESKTOP, '审计底稿_TB勾稽复核报告.csv')
    with open(outp, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow(['【A】资产负债表类审定表 vs TB 期末余额(signed)'])
        w.writerow(['文件夹', '文件', '审定表sheet', '科目编码', '年度', '报告期末(签名)', 'TB期末(签名)', '结论'])
        for r in rows_bs:
            w.writerow(r)
        w.writerow([])
        w.writerow(['【B】利润表 GL自然方发生额(已应用蓝字规则) vs TB自然方发生额'])
        w.writerow(['文件夹', '科目', '主体', '年度', '二级', 'GL自然方', 'TB自然方', '结论'])
        for r in rows_pl:
            w.writerow(r)

    bs = [r for r in rows_bs if len(r) >= 8 and r[7] != 'NO_HEADER' and r[7] != 'NO_COL']
    bs_ok = [r for r in bs if r[7] == 'OK']
    pl = [r for r in rows_pl if len(r) >= 8 and r[7] not in ('NO_TB',)]
    pl_ok = [r for r in pl if r[7] == 'OK']
    pl_ytd = [r for r in pl if r[7].startswith('YTD')]
    pl_diff = [r for r in pl if r[7].startswith('DIFF')]
    print('=' * 70)
    print('审计底稿 ↔ TB 勾稽复核')
    print('  (A) 资产负债表审定表: %d/%d 行一致' % (len(bs_ok), len(bs)))
    print('  (B) 利润表 GL自然==TB自然: %d/%d 行一致；YTD不可比 %d；实差 %d'
          % (len(pl_ok), len(pl), len(pl_ytd), len(pl_diff)))
    print('  报告: %s' % outp)
    if bs_diff := [r for r in bs if r[7].startswith('DIFF')]:
        print('\n  [A 差异]')
        for r in bs_diff[:40]:
            print('   ', r)
    if pl_diff:
        print('\n  [B 实差(非YTD)]')
        for r in pl_diff[:40]:
            print('   ', r)


if __name__ == "__main__":
    main()
