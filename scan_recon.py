# -*- coding: utf-8 -*-
"""全科目底稿勾稽扫描：对 {dir} 下所有 *审计底稿*.xlsx，
对比 审定表期末 vs 明细表合计期末 vs 附注合计，输出不一致科目。
用法：python scan_recon.py <底稿目录> [--detail]
"""
import sys, os, glob, re
import openpyxl

def find_end_col(ws, n_hdr=6):
    """在表头行（前 n_hdr 行）找期末列。按优先级匹配：
    期末审定数 > 期末余额(人民币) > 期末余额 > 期末未审 > 期末数 > 期末。
    返回列号(1-based)或 None。"""
    prio = ['期末审定数', '期末余额(人民币)', '期末余额', '期末未审', '期末数', '期末']
    for p in prio:
        for r in range(1, n_hdr + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is None:
                    continue
                s = str(v).replace(' ', '')
                if s == p or (p == '期末' and '期末' in s and '期初' not in s):
                    return c
    return None

def find_sum_end(ws, end_col):
    """找合计行（第1/2列含 合计/小计/年度合计）的期末值。"""
    for r in range(1, ws.max_row + 1):
        v1 = ws.cell(r, 1).value
        v2 = ws.cell(r, 2).value
        if v1 is not None and ('合计' in str(v1) or '小计' in str(v1)) or (v2 is not None and ('合计' in str(v2) or '小计' in str(v2))):
            v = ws.cell(r, end_col).value
            if isinstance(v, (int, float)):
                return v
    # 无合计行：取数据区（表头后）最后数值
    vals = []
    for r in range(7, ws.max_row + 1):
        v = ws.cell(r, end_col).value
        if isinstance(v, (int, float)):
            vals.append((r, v))
    return vals[-1][1] if vals else None

def find_max_end(ws, end_col):
    """审定表：取数据区该列绝对值最大的值（期末数一般最大）。"""
    best = None
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, end_col).value
        if isinstance(v, (int, float)) and abs(v) > 0.005:
            if best is None or abs(v) > abs(best):
                best = v
    return best

def scan_one(path, detail=False):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = wb.sheetnames
    aud_end = det_end = note_end = None
    aud_sheet = det_sheet = note_sheet = None
    for sn in sheets:
        ws = wb[sn]
        if '审定表' in sn and aud_end is None:
            ec = find_end_col(ws)
            if ec:
                aud_end = find_max_end(ws, ec)
                aud_sheet = sn
        elif '明细表' in sn and '附注' not in sn and det_end is None:
            ec = find_end_col(ws)
            if ec:
                det_end = find_sum_end(ws, ec)
                det_sheet = sn
        elif '附注' in sn and note_end is None:
            ec = find_end_col(ws)
            if ec:
                note_end = find_sum_end(ws, ec)
                note_sheet = sn
    wb.close()
    diffs = []
    if aud_end is not None and det_end is not None:
        d = det_end - aud_end
        if abs(d) > 50:
            diffs.append(f'明细vs审定 {d:,.2f}')
    if aud_end is not None and note_end is not None:
        d = note_end - aud_end
        if abs(d) > 50:
            diffs.append(f'附注vs审定 {d:,.2f}')
    if diffs or detail:
        return os.path.basename(path), aud_end, det_end, note_end, diffs, (aud_sheet, det_sheet, note_sheet)
    return None

def main():
    d = sys.argv[1] if len(sys.argv) > 1 else r'D:/底稿测试/AH/tmp_1010full'
    detail = '--detail' in sys.argv
    files = sorted(glob.glob(os.path.join(d, '*审计底稿_*.xlsx')))
    if not files:
        files = sorted(glob.glob(os.path.join(d, '*.xlsx')))
    print(f'扫描 {len(files)} 个底稿（{d}）')
    print('=' * 90)
    n_diff = 0
    for f in files:
        if os.path.basename(f).startswith('~$'):
            continue
        try:
            res = scan_one(f, detail)
        except Exception as e:
            print(f'{os.path.basename(f)}: ERR {e}')
            continue
        if res:
            name, aud, det, note, diffs, sheets = res
            n_diff += 1
            print(f'{name}')
            print(f'  审定表[{sheets[0]}]: {aud if aud is None else f"{aud:,.2f}"} | 明细[{sheets[1]}]: {det if det is None else f"{det:,.2f}"} | 附注[{sheets[2]}]: {note if note is None else f"{note:,.2f}"}')
            for dd in diffs:
                print(f'    ⚠ {dd}')
    print('=' * 90)
    print(f'差异底稿数: {n_diff}/{len(files)}')

if __name__ == '__main__':
    main()
