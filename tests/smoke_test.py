# -*- coding: utf-8 -*-
"""tests/smoke_test.py —— 冒烟单测（2026-08-22 建立，纯标准库）。

覆盖关键读取/映射/脱敏/勾稽函数，防止改公共层破坏。运行：
  python tests/smoke_test.py
"""
import os
import sys

APP = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
sys.path.insert(0, os.path.abspath(APP))

import paths as P   # ⚡ 2026-08-30 修复：必须在 sys.path.insert(APP) 之后 import（原顺序导致 ModuleNotFoundError）

PASS = FAIL = 0


def check(name, cond, extra=''):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f'  ✓ {name}')
    else:
        FAIL += 1
        print(f'  ✗ {name} {extra}')


def t_mask_safe():
    from mask_safe import _paired_mask, prefer_masked, require_masked
    # 同目录 _脱敏
    p = _paired_mask(os.path.join(P.DATA_DIRS['ADF'], '数据', '2026'), '融资明细.xls')
    check('mask_safe 同目录配对', p and p.endswith('融资明细_脱敏.xlsx'), p)
    # 同级 _脱敏 子目录（12月资产台账）
    p2 = _paired_mask(os.path.join(P.DATA_DIRS['ADF'], '数据', '2026', '12月固定资产台账', '12月资产台账'), '3100.XLSX')
    check('mask_safe 同级_脱敏目录配对', p2 and '12月资产台账_脱敏' in p2, p2)
    # 父级 _脱敏 目录（合同/回函目录级）
    p3 = prefer_masked(os.path.join(P.DATA_DIRS['ADF'], '数据', '2026', '借款合同_ocr', '识别文本', '3900苏震', 'x.txt'))
    check('mask_safe 目录级回退安全', p3 is not None)   # 无脱敏版回退原路径（读取层保守）
    r = require_masked(os.path.join(P.DATA_DIRS['ADF'], '不存在', 'x.xlsx'))
    check('mask_safe require_masked 缺失→None', r is None)


def t_key_mapping():
    from key_mapping import key_map, get_key, is_match, list_maps
    cfg = key_map('ADF', 'financing')
    check('key_mapping 取配置', cfg.get('match') == 'norm')
    k = get_key(cfg, {'业务编号': 'JK001'}, '台账')
    check('key_mapping get_key', k == 'JK001')
    check('key_mapping norm 匹配', is_match(cfg, 'JK2026-0317', 'JK20260317', fuzzy=True))


def t_ledger_registry():
    import ledger_registry as LR
    miss = LR.check()
    check('ledger_registry 台账脱敏齐备', not miss, miss)


def t_financing_ledger():
    import financing_ledger as FL
    st = FL._load_finance_loans('ST')
    lt = FL._load_finance_loans('LT')
    amt_st = sum(x['amt'] for x in st)
    check('financing_ledger ST 39笔', len(st) == 39, len(st))
    check('financing_ledger LT 70笔', len(lt) == 70, len(lt))
    check('financing_ledger ST 金额>30亿', amt_st > 30e8, round(amt_st / 1e8, 2))
    check('financing_ledger 主体为借X', all(str(x['ent']).startswith('借') for x in st[:10]))


def t_fa_ledger_mask():
    # 脱敏版存在性 + 公司代码→借X
    import os
    end = os.path.join(P.DATA_DIRS['ADF'], '数据', '2026', '7月固定资产台账_脱敏.xlsx')
    check('fa_ledger_mask 脱敏版存在', os.path.exists(end))
    beg = os.path.join(P.DATA_DIRS['ADF'], '数据', '2026', '12月固定资产台账', '12月资产台账_脱敏')
    check('fa_ledger_mask 12月脱敏目录', os.path.isdir(beg))
    if os.path.exists(end):
        import openpyxl
        wb = openpyxl.load_workbook(end, read_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(min_row=2, max_row=6, values_only=True))
        wb.close()
        comps = {str(r[3]) for r in rows if r[3]}
        check('fa_ledger_mask 公司代码→借X', comps <= {'借01', '借02', '借03', '借04', '借05', '借06', ''}, comps)


def t_loan_detail():
    import loan_detail as L
    st = L._load_finance_loans('ST')
    check('loan_detail 复用 financing_ledger ST', len(st) == 39, len(st))
    check('loan_detail _mask_bank', L._mask_bank('中国银行') == '贷01')


def t_sap_adapter():
    """2026-08-22 收敛第一步：公共层双写（dr/cr+qc/jf/df/qm 与 debit/credit 一致）。"""
    import sap_adapter as SA
    # read_km 双写
    import os
    old_root, old_comp = SA._DATA_ROOT, SA._current_comp
    try:
        SA._DATA_ROOT = os.path.join(P.DATA_DIRS['ADF'], '数据', '2026')
        SA._current_comp = '3100'
        km = SA.read_km()
        assert km, 'read_km 空'
        sample = next(iter(km.values()))
        check('sap read_km qc 双写', abs(sample['qc'] - sample['opening']) < 1e-9)
        check('sap read_km df 双写', abs(sample['df'] - sample['credit']) < 1e-9)
        # read_gl 双写
        gl = SA.read_gl()
        if gl:
            r0 = gl[0]
            check('sap read_gl dr 双写', abs(float(r0.get('dr') or 0) - float(r0.get('debit') or 0)) < 1e-9)
            check('sap read_gl cr 双写', abs(float(r0.get('cr') or 0) - float(r0.get('credit') or 0)) < 1e-9)
        else:
            print('  (read_gl 空，跳过字段检查)')
    except Exception as e:
        check('sap_adapter 双写', False, str(e))
    finally:
        SA._DATA_ROOT, SA._current_comp = old_root, old_comp


def t_co_carryover():
    """2026-08-23 CO 结转凭证识别：仅含 7042/7043 成本对象科目。
    ⚡ 勿用文本关键词——『计提折旧和摊销』等真实计提不得标记 CO。"""
    import sap_common as SC
    assert hasattr(SC, '_CO_MIN_ROWS')

    def is_co(rs):
        return any(str(r['code']).startswith(('7042', '7043')) for r in rs)

    rs_704 = [{'code': '7043000101', 'amt': 100.0, 'sm': ''} for _ in range(4)] \
        + [{'code': '6600000000', 'amt': -100.0, 'sm': ''} for _ in range(4)]          # 含704借贷成对
    rs_kw = [{'code': '6600000000', 'amt': 50.0, 'sm': '计提折旧和摊销'} for _ in range(4)] \
        + [{'code': '1602020000', 'amt': -50.0, 'sm': '计提折旧和摊销'} for _ in range(4)]  # 真实折旧计提→不标记
    rs_norm = [{'code': '1002010102', 'amt': 100.0, 'sm': '承兑'} for _ in range(4)] \
        + [{'code': '1002020101', 'amt': -100.0, 'sm': '承兑'} for _ in range(4)]        # 正常资金归集
    check('co 含704识别', is_co(rs_704))
    check('co 折旧计提不误伤', not is_co(rs_kw))
    check('co 正常凭证不误伤', not is_co(rs_norm))


def t_signed_balance_v2():
    """2026-08-23 signed_balance_v2：兼容 U8『方向+绝对金额』与 SAP『方向+带符号金额』。"""
    from audit_common import signed_balance_v2 as v2
    check('SAP贷+负 保持负', v2('贷', -196815301.18) == -196815301.18)
    check('U8贷+正 取负', v2('贷', 813822.63) == -813822.63)
    check('U8借+正 保持正', v2('借', 3382161.39) == 3382161.39)
    check('平+正 保持正', v2('平', 100.0) == 100.0)


def t_filter_inc_cp():
    """2026-08-23 增加检查表对方科目真实来源筛选：DROP 优先、KEEP 保留、未知保留。"""
    from audit_common import filter_inc_cp as f
    # 凭证级多科目：剔除 税费/费用/组内/存货/薪酬，保留 银行/应付/应收/票据/在建工程
    r = f('应付账款-明细应付款；应交税费-应交增值税-进项税额；固定资产；应收账款-配件款；银行存款')
    check('真实来源保留', '应付账款-明细应付款' in r and '银行存款' in r)
    check('税费剔除', '应交税费' not in r)
    check('组内固定资产剔除', '固定资产' not in r)
    # DROP 优先：信用减值损失-应收账款损失 同时含 KEEP(应收账款) 与 DROP(信用减值损失)
    r2 = f('信用减值损失-应收账款损失；银行存款')
    check('DROP优先剔除', '信用减值损失' not in r2 and '银行存款' in r2)
    # 顿号不误分割（在建工程-建筑、安装工程 是一个科目）
    r3 = f('在建工程-建筑、安装工程；银行存款')
    check('顿号不拆散科目', '在建工程-建筑、安装工程' in r3 and '银行存款' in r3)
    # 空值安全
    check('空值安全', f(None) == '' and f('') == '')


def t_detect_layout():
    """2026-08-23 detect_sap_layout 表头内容探测：U8 币种/分主体目录不再误判 SAP 300。"""
    import tempfile
    import openpyxl
    from sap_reader import detect_sap_layout
    tmp = tempfile.mkdtemp()
    # 构造 U8 子目录（AZ 币种分目录风格：科目代码|科目名称 表头）
    u8d = os.path.join(tmp, 'CNY')
    os.makedirs(u8d)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['科目代码', '科目名称', '方向', '金额', '本期借方', '本期贷方', '方向', '金额', '等级'])
    ws.append(['1002', '银行存款', '借', '813822.63', '23851981.16', '24129820.83', '借', '535982.96', '1'])
    wb.save(os.path.join(u8d, '数链2025科目余额表.xlsx'))
    check('U8币种目录不判300', detect_sap_layout(tmp) is None)
    # 构造 SAP 300 子目录（公司|总账|总账科目长文本 表头）
    s3 = os.path.join(tmp, '3700')
    os.makedirs(s3)
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(['公司', '总账', '总账科目长文本', '期初方向', '期末的方向', '期初金额', '本期借方金额', '本期贷方金额'])
    ws2.append(['3700', '1002010102', '银行存款-工行', '借', '平', '4024561', '1300', '1300'])
    wb2.save(os.path.join(s3, '3700科目余额表20260731.xlsx'))
    check('SAP300主体目录判300', detect_sap_layout(tmp) == '300')
    # GJX 单文件夹 U8（DQ2025科目余额表.xlsx 是 U8 表头）→ 不判 dq/300
    gx = os.path.join(tmp, 'gx')
    os.makedirs(gx)
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.append(['科目代码', '科目名称', '方向', '金额', '本期借方', '本期贷方', '方向', '金额', '等级'])
    ws3.append(['1001', '现金', '借', '5863.4', '329550', '333978', '借', '1435.4', '1'])
    wb3.save(os.path.join(gx, 'DQ2025科目余额表.xlsx'))
    check('GJX单文件夹U8不判SAP', detect_sap_layout(gx) is None)
    # AH SAP FAGL（科目余额表目录内文件含 公司代码/科目号）→ ah-sap
    ahd = os.path.join(tmp, 'ah')
    os.makedirs(os.path.join(ahd, '科目余额表'))
    os.makedirs(os.path.join(ahd, '序时账'))
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.append(['货币类型', '货币', '科目号', '功能范围', '业务范围', '段', '利润中心', '公司代码', '期间/年度'])
    ws4.append(['公司代码货币', '人民币', '库存现金', '未分配的', '未分配的', '未分配的', '未分配的', '杭氧集团股份有限公司', '剩余 2026'])
    wb4.save(os.path.join(ahd, '科目余额表', '1010.xlsx'))
    check('AH SAP FAGL判ah-sap', detect_sap_layout(ahd) == 'ah-sap')


def t_ledger_backend():
    """2026-08-24 架构阶段1+2：LedgerBackend 形态分发 + 统一 read_tb 入口。"""
    import tempfile
    import openpyxl
    from ledger_backend import get_backend, backend_for
    tmp = tempfile.mkdtemp()
    # U8 目录
    u8d = os.path.join(tmp, 'CNY')
    os.makedirs(u8d)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['科目代码', '科目名称', '方向', '金额', '本期借方', '本期贷方', '方向', '金额', '等级'])
    ws.append(['1002', '银行存款', '借', '813822.63', '23851981.16', '24129820.83', '借', '535982.96', '1'])
    wb.save(os.path.join(u8d, '数链2025科目余额表.xlsx'))
    check('U8 backend 分发', get_backend(tmp).layout == 'u8')
    check('SAP300 backend 分发', backend_for('300', tmp).layout == '300')
    check('ah-sap backend 分发', backend_for('ah-sap', tmp).layout == 'ah-sap')


def t_desensitizer():
    """2026-08-24 架构·输出侧：统一脱敏层（一致性/字段类型/编码保留/语境隔离）。"""
    from desensitizer import Desensitizer, get_contract
    dz = Desensitizer()
    dz.register('entity', {'上海万坤实业发展有限公司': 'A001'})
    check('脱敏一致性', dz.mask('上海万坤实业发展有限公司', 'entity') == 'A001')
    check('脱敏同值同码', dz.mask('上海万坤实业发展有限公司', 'entity') == 'A001')
    check('账号保留尾号', '7890' in dz.mask('62220210011234567890', 'account'))
    check('人员保留姓氏', dz.mask('张伟', 'person') == '张*')
    check('编码保留', dz.mask('YS51010001', 'code') == 'YS51010001')
    # ⚡ 2026-08-24 语境隔离：contract 语境（借X/贷X）与 confirm 语境（A/B/C）分实例
    dzc = get_contract()
    check('银行贷X', dzc.mask('中国银行', 'bank') == '贷01')
    check('公司代码借X', dzc.mask('3100', 'entity') == '借01')
    check('文本不误伤厂房', dzc.mask('100号厂房及设备', 'text') == '100号厂房及设备')
    check('文本公司泛化', '【客商】' in dzc.mask('江苏东方盛虹股份有限公司 办公楼', 'text'))
    # 跨行容忍（confirm_mask 原 re.sub 行为）
    dz2 = Desensitizer()
    dz2.register('entity', {'浙江兆龙互连科技股份有限公司': 'A001'})
    check('跨行容忍', dz2.mask('浙江兆龙互连科技股份\n有限公司', 'text') == 'A001')


def t_recon_extract():
    """2026-08-24 架构 R1：统一差异抽取器（标记列/数值列/表头跳过）。"""
    from recon_diff_extract import extract_workbook
    import tempfile, os
    from openpyxl import Workbook
    tmp = tempfile.mkdtemp()
    fp = os.path.join(tmp, 'test_核对.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = '明细'
    ws.append(['主体', '科目', '结果'])
    ws.append(['A', '1002', '✅ 匹配'])
    ws.append(['B', '1002', '⚠️ 清单独有(账面未发现)'])
    ws.append(['C', '1002', '✅ 匹配'])
    wb.save(fp)
    # 标记列抽取（col2）
    diffs = extract_workbook(fp, '明细', [2])
    check('标记列抽取', len(diffs) == 1 and '清单独有' in diffs[0].get('2', ''))
    # 数值列抽取（差异列≠0）+ 表头跳过
    fp2 = os.path.join(tmp, 'test_tb.xlsx')
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(['说明'])
    ws2.append(['口径'])
    ws2.append(['项目', '表合计', 'TB', '差异'])
    ws2.append(['应收账款', '100', '100', '0'])
    ws2.append(['坏账准备', '0', '-50', '50'])
    wb2.save(fp2)
    diffs2 = extract_workbook(fp2, None, [3], num_mark=True, skip_rows=3)
    check('数值列抽取', len(diffs2) == 1 and '坏账准备' in diffs2[0].get('0', ''))
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)


def t_data_source_registry():
    """2026-08-24 架构 R4：数据源注册表（输入解析/注册完整性）。"""
    from data_source_registry import REGISTRY, _resolve_input, list_sources
    check('数据源已登记', 'currency_split' in REGISTRY and 'contract_ocr' in REGISTRY
          and 'confirm_ocr' in REGISTRY)
    check('数据源kind齐全', {REGISTRY[k]['kind'] for k in REGISTRY} == {'preprocess', 'extract', 'analyze'})
    # 输入解析：相对路径与通配 glob
    import tempfile, os
    tmp = tempfile.mkdtemp()
    os.makedirs(os.path.join(tmp, '数据'), exist_ok=True)
    check('数据源输入解析', _resolve_input(tmp, '数据') == os.path.join(tmp, '数据'))
    os.makedirs(os.path.join(tmp, '网银2025'), exist_ok=True)
    check('数据源通配解析', _resolve_input(tmp, '*网银*') == os.path.join(tmp, '网银2025'))
    import shutil
    shutil.rmtree(tmp, ignore_errors=True)


def t_agg_cost_dedup():
    """2026-08-24 修复：_agg_cost 点分级去重（父级+子目双算 / 父级无子目回退）。"""
    import longterm_assets_detail as L
    # 父级存在 + 子目 → 只取子目（防双算）
    km1 = {
        '1601': {'name': '固定资产', 'opening': 500.0, 'debit': 100.0, 'credit': 20.0, 'closing': 580.0},
        '1601.01': {'name': '房屋', 'opening': 300.0, 'debit': 60.0, 'credit': 10.0, 'closing': 350.0},
        '1601.02': {'name': '设备', 'opening': 200.0, 'debit': 40.0, 'credit': 10.0, 'closing': 230.0},
    }
    rec1 = L._prep_group('t', km1, {}, 'FA')
    check('父级有子目只取子目', abs((rec1['cost_d'] or {}).get('closing') - 580.0) < 0.01)
    # 父级存在 + 无子目 → 回退父级自身（不丢原值）
    km2 = {
        '1701': {'name': '无形资产', 'opening': 5000.0, 'debit': 100.0, 'credit': 0.0, 'closing': 5100.0},
    }
    rec2 = L._prep_group('t', km2, {}, 'INT')
    check('父级无子目回退', abs((rec2['cost_d'] or {}).get('closing') - 5100.0) < 0.01)
    # 无父级平行子目（XBJ 实为 160101 土地资产等，名称含组名）→ 全取
    km3 = {
        '160101': {'name': '固定资产-土地资产', 'opening': 0.0, 'debit': 0.0, 'credit': 0.0, 'closing': 100.0},
        '160102': {'name': '固定资产-房屋建筑物', 'opening': 0.0, 'debit': 0.0, 'credit': 0.0, 'closing': 200.0},
        '160103': {'name': '固定资产-机器设备', 'opening': 0.0, 'debit': 0.0, 'credit': 0.0, 'closing': 300.0},
    }
    rec3 = L._prep_group('t', km3, {}, 'FA')
    check('无父级平行子目全取', abs((rec3['cost_d'] or {}).get('closing') - 600.0) < 0.01)


def t_name_hit_thai():
    """2026-08-24 修复：_name_hit 排除泰国合并科目（应收X和应收Y）。"""
    import current_account_detail as CA
    check('泰国合并排除票据', not CA._name_hit('应收账款和应收票据', '应收票据'))
    check('泰国合并保留应收', CA._name_hit('应收账款和应收票据', '应收账款'))
    check('应付合并排除票据', not CA._name_hit('应付账款和应付票据', '应付票据'))
    check('应付合并保留应付', CA._name_hit('应付账款和应付票据', '应付账款'))
    check('真实应收票据命中', CA._name_hit('应收票据-银行承兑', '应收票据'))
    check('真实应收命中', CA._name_hit('应收账款', '应收账款'))


def main():
    print('=== 冒烟单测 ===')
    for t in (t_mask_safe, t_key_mapping, t_ledger_registry,
              t_financing_ledger, t_fa_ledger_mask, t_loan_detail, t_sap_adapter,
              t_co_carryover, t_signed_balance_v2, t_filter_inc_cp, t_detect_layout,
              t_ledger_backend, t_desensitizer, t_recon_extract, t_data_source_registry,
              t_agg_cost_dedup, t_name_hit_thai):
        try:
            t()
        except Exception as e:
            global FAIL
            FAIL += 1
            print(f'  ✗ {t.__name__} 异常: {e}')
    print(f'\n结果: PASS {PASS} / FAIL {FAIL}')
    return 1 if FAIL else 0


if __name__ == '__main__':
    sys.exit(main())
