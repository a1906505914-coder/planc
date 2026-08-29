# -*- coding: utf-8 -*-
"""底稿完整性检查（2026-08-29 用户核心诉求：给结果时不想逐张点开底稿查空白/缺失）。

把散落的检查整合成一张"完整性总表"，一眼看出哪些底稿有问题：
  1. 关键表（审定表/明细表）空表判定（TB 有数据但表空 = ERROR，复用 sheet_gap_check）
  2. 数值校验（合计/小计与明细不符等 = ERROR/WARN，复用 audit_checker.check_file）
  3. 关键表行数统计（内容量级，供人工判断"该有多少数据"）
产出：完整性总表（每底稿一行）+ 异常明细 + 状态汇总。

用法：
    python wp_completeness_check.py                       # 默认扫 AH 三集团
    python wp_completeness_check.py --dir <底稿目录> --data <数据根目录>
    python wp_completeness_check.py --report <输出.md>

返回码：0=全部完整 | 1=有 ERROR | 2=有 WARN。
"""
import glob
import os
import re
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import audit_checker
import sheet_gap_check

AH_GROUP_ROOT = r'd:/底稿测试/AH/底稿/2026/集团'
AH_DATA = r'd:/底稿测试/AH/数据/2026'

# 关键表（决定底稿"是否有内容"）：含这些关键词的 sheet 空=异常
_KEY_SHEET_KW = ('审定表', '明细表', '附注汇总', '余额表', '台账')


def _key_sheets(ws_list):
    return [s for s in ws_list if any(k in s for k in _KEY_SHEET_KW)]


def _count_numeric_rows(ws):
    """关键表数值行数（剔除表头/合计/小计/说明行后的非零数值行数）。"""
    n = 0
    try:
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            a = str(row[0]).strip()
            if not a or a.startswith(('勾稽', '注', '说明', '口径', '合计', '小计', '总计')):
                continue
            if any(k in a for k in ('核算主体', '科目', '项目', '期间', '年度', '序号', '账户', '对方科目', '摘要', '日期')):
                continue
            for c in row[1:]:
                if isinstance(c, (int, float)) and abs(c) > 0.005:
                    n += 1
                    break
    except Exception:
        pass
    return n


def check_one(fp, tb_has):
    """单底稿完整性。返回 (status, info_dict)。status: OK/WARN/ERROR。"""
    base = os.path.basename(fp)
    err_msgs, warn_msgs = [], []
    ks, row_counts = [], {}
    # 1) 数值校验（audit_checker）
    try:
        for lvl, msg in audit_checker.check_file(fp):
            if lvl == 'ERROR':
                err_msgs.append(msg)
            else:
                warn_msgs.append(msg)
    except Exception as _ex:
        warn_msgs.append(f'audit_checker 异常: {_ex}')
    # 2) 关键表空表 + 行数统计
    try:
        import openpyxl
        wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
        ws_list = wb.sheetnames
        ks = _key_sheets(ws_list)
        empty_ks = []
        for sn in ks:
            nrow = _count_numeric_rows(wb[sn])
            row_counts[sn] = nrow
            if nrow == 0:
                empty_ks.append(sn)
        wb.close()
        # TB 有数据但关键表全空 → ERROR
        tb_flag = tb_has.get(base, True)
        if tb_flag and ks and len(empty_ks) == len(ks):
            err_msgs.append('关键表全空（TB 有数据）')
        elif not tb_flag and ks and len(empty_ks) == len(ks):
            warn_msgs.append('关键表空（TB 无数据，属正常跳过）')
        elif empty_ks:
            warn_msgs.append(f'部分关键表空：{";".join(empty_ks)[:60]}')
    except Exception as _ex:
        warn_msgs.append(f'打开失败: {_ex}')
    if err_msgs:
        status = 'ERROR'
    elif warn_msgs:
        status = 'WARN'
    else:
        status = 'OK'
    return status, {
        'file': base, 'err': err_msgs, 'warn': warn_msgs,
        'n_key': len(ks), 'row_counts': row_counts,
    }


def scan(root, data_root):
    """扫描 root 下全部集团目录/单目录，返回 (rows, detail)。"""
    dirs = []
    if os.path.isdir(root):
        sub = [os.path.join(root, d) for d in sorted(os.listdir(root))
               if os.path.isdir(os.path.join(root, d))]
        if sub and all(re.fullmatch(r'\d{4}', os.path.basename(d)) for d in sub):
            dirs = sub            # 集团根：1010/1357/2468 子目录
        else:
            dirs = [root]         # 单目录
    # ⚡⚡ 2026-08-29 P0 修复：build_tb_map 的 folder 必须含 xlsx 文件（os.listdir 只见文件）。
    #   集团根目录下只有子目录 → 空 map → 全部默认"TB 有数据"→ 无数据科目被误报 ERROR。
    #   改为对每个含文件的子目录分别构建再合并。
    tb_map = {}
    for d in dirs:
        try:
            tb_map.update(sheet_gap_check.build_tb_map(data_root, d))
        except Exception:
            pass
    files = []
    for d in dirs:
        files += [(d, f) for f in sorted(glob.glob(os.path.join(d, '*审计底稿*.xlsx')))
                  if not os.path.basename(f).startswith('~$')]
    rows, detail = [], []
    t0 = time.time()
    for d, fp in files:
        g = os.path.basename(d)
        status, info = check_one(fp, tb_map)
        rows.append((g, info['file'], len(info.get('row_counts', {})),
                     sum(1 for v in info.get('row_counts', {}).values() if v == 0),
                     len(info['err']), len(info['warn']), status))
        if status != 'OK':
            detail.append((g, info['file'], status, info['err'], info['warn']))
        print(f'  {"✅" if status=="OK" else ("⚠️" if status=="WARN" else "❌")} {g}/{info["file"]} '
              f'关键表{len(info.get("row_counts", {}))} 空{sum(1 for v in info.get("row_counts", {}).values() if v==0)} '
              f'E{len(info["err"])}/W{len(info["warn"])}（{time.time()-t0:.0f}s）', flush=True)
    return rows, detail


def write_report(rows, detail, out_path, elapsed):
    n_ok = sum(1 for r in rows if r[6] == 'OK')
    n_warn = sum(1 for r in rows if r[6] == 'WARN')
    n_err = sum(1 for r in rows if r[6] == 'ERROR')
    L = []
    L.append('# 底稿完整性检查报告')
    L.append(f'- 时间：{time.strftime("%Y-%m-%d %H:%M:%S")}（扫描 {elapsed:.0f}s）')
    L.append(f'- 底稿数：{len(rows)}  ✅完整 {n_ok} / ⚠️WARN {n_warn} / ❌ERROR {n_err}')
    L.append('')
    L.append('| 集团 | 底稿 | 关键表 | 空关键表 | ERROR | WARN | 状态 |')
    L.append('|---|---|---|---|---|---|---|')
    for g, f, nk, ne, ne2, nw, st in rows:
        tag = {'OK': '✅', 'WARN': '⚠️', 'ERROR': '❌'}.get(st, '?')
        L.append(f'| {g} | {f} | {nk} | {ne} | {ne2} | {nw} | {tag} |')
    if detail:
        L.append('')
        L.append('## 异常明细')
        for g, f, st, errs, warns in detail:
            L.append(f'### {g}/{f}（{st}）')
            for m in (errs or [])[:6]:
                L.append(f'- ❌ {m}')
            for m in (warns or [])[:6]:
                L.append(f'- ⚠️ {m}')
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(L))
    return n_err, n_warn


def main():
    argv = sys.argv[1:]
    root = AH_GROUP_ROOT
    data = AH_DATA
    if '--dir' in argv:
        root = argv[argv.index('--dir') + 1]
    if '--data' in argv:
        data = argv[argv.index('--data') + 1]
    out = os.path.join(HERE, 'completeness_reports', f'完整性检查_{time.strftime("%Y%m%d_%H%M%S")}.md')
    if '--report' in argv:
        out = os.path.abspath(argv[argv.index('--report') + 1])
    print(f'═════ 底稿完整性检查：{root} ═════')
    t0 = time.time()
    rows, detail = scan(root, data)
    n_err, n_warn = write_report(rows, detail, out, time.time() - t0)
    print(f'\n汇总：{len(rows)} 份底稿 / ❌ERROR {n_err} / ⚠️WARN {n_warn}')
    print(f'报告已写：{out}')
    if n_err:
        print('⛔ 存在 ERROR，请处理后再交付（返回码 1）')
        return 1
    if n_warn:
        print('⚠️ 存在 WARN，建议确认（返回码 2）')
        return 2
    print('✅ 全部底稿完整（返回码 0）')
    return 0


if __name__ == '__main__':
    sys.exit(main())
